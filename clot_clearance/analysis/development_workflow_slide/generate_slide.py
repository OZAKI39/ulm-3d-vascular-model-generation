"""Build exactly one editable workflow slide from archived evidence only."""
import json,re,hashlib,os,shutil,subprocess,sys,math
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR,MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from PIL import ImageFont
from slide_config import SLIDE_CONFIG as C
from make_insets import make,ROOT,V,HERE
NAME='development_workflow_current_progress'
FONT_DIR=Path('/usr/share/fonts/truetype/dejavu')
TEXT_RECORDS=[]

def rgb(hex):return RGBColor.from_string(hex)
def wrap(text,w,size,bold=False):
 font=ImageFont.truetype(str(FONT_DIR/('DejaVuSans-Bold.ttf' if bold else 'DejaVuSans.ttf')),round(size*4))
 lines=[]
 for para in text.split('\n'):
  line=''
  for word in para.split():
   test=(line+' '+word).strip()
   if line and font.getlength(test)>max(1,w*72-3)*4:lines.append(line);line=word
   else:line=test
  lines.append(line)
 return lines

def text(slide,s,x,y,w,h,size=13,color='243746',bold=False,align=PP_ALIGN.LEFT):
 assert not re.search('[\u3400-\u9fff]',s)
 lines=wrap(s,w,size,bold)
 required=len(lines)*size*1.13/72
 if required>h+.035:raise ValueError(f'Text exceeds box: {s!r}; {required:.3f}>{h}')
 sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));sh.name='Editable text: '+s.replace('\n',' ')[:70]
 tf=sh.text_frame;tf.clear();tf.auto_size=MSO_AUTO_SIZE.NONE;tf.word_wrap=False;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.vertical_anchor=MSO_ANCHOR.TOP
 for i,line in enumerate(lines):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.alignment=align;p.space_before=p.space_after=Pt(0);p.line_spacing=Pt(size*1.13)
  p.font.name=C['font'];p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
 TEXT_RECORDS.append(dict(text=s,x=x,y=y,w=w,h=h,size_pt=size,line_count=len(lines)))
 return sh

def shape(slide,kind,x,y,w,h,fill,line=None,lw=.7,name='Editable shape'):
 s=slide.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h));s.name=name
 if fill:s.fill.solid();s.fill.fore_color.rgb=rgb(fill)
 else:s.fill.background()
 if line:s.line.color.rgb=rgb(line);s.line.width=Pt(lw)
 else:s.line.fill.background()
 s._element.spPr.append(OxmlElement('a:effectLst'))
 for ref in s._element.xpath('./p:style/a:effectRef'):ref.set('idx','0')
 return s

def line(slide,x1,y1,x2,y2,color='687580',width=.7):
 s=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));s.line.color.rgb=rgb(color);s.line.width=Pt(width)
 s._element.spPr.append(OxmlElement('a:effectLst'))
 for ref in s._element.xpath('./p:style/a:effectRef'):ref.set('idx','0')
 return s

def clot(slide,x,y,w):
 text(slide,'TARGET APPLICATION',x,y,w,.23,11.5,'687580',True)
 # A geometric schematic only; no simulated contours or crack geometry.
 yy=y+.47
 shape(slide,MSO_SHAPE.RECTANGLE,x+.08,yy,w-.16,1.14,'FFFFFF','BCC6CD',.7,'Vessel lumen')
 shape(slide,MSO_SHAPE.RECTANGLE,x+.08,yy+.85,w-.16,.29,'EEF3F6',None,name='Near-wall region')
 shape(slide,MSO_SHAPE.RIGHT_ARROW,x+.42,yy+.25,w-.86,.14,'4D738A',None,name='Flow direction arrow')
 text(slide,'Flow',x+.40,yy+.02,1.0,.24,12,'4D738A')
 points=[(Inches(x+1.56+.56*math.cos(t)),Inches(yy+1.07-.45*math.sin(t))) for t in [math.pi*(1-i/24) for i in range(25)]]
 builder=slide.shapes.build_freeform(points[0][0],points[0][1]);builder.add_line_segments(points[1:],close=True);mass=builder.convert_to_shape();mass.name='Attached clot native vector schematic';mass.fill.solid();mass.fill.fore_color.rgb=rgb('A6BDC9');mass.line.color.rgb=rgb('527384');mass.line.width=Pt(.8);mass._element.spPr.append(OxmlElement('a:effectLst'))
 for ref in mass._element.xpath('./p:style/a:effectRef'):ref.set('idx','0')
 shape(slide,MSO_SHAPE.RECTANGLE,x+.08,yy+1.07,w-.16,.18,'BDC7CD',None,name='Vessel wall')
 text(slide,'Attached clot',x+1.66,yy+.45,1.5,.24,11.5,'344A58')
 line(slide,x+1.73,yy+.70,x+1.63,yy+.79)
 text(slide,'Near-wall\nregion',x+.13,yy+.61,.90,.40,11,'5A6B76')
 text(slide,'Vessel wall',x+.14,yy+1.22,1.55,.23,11.5,'5A6B76')
 text(slide,C['future_target'],x,yy+1.51,w,.55,11.5,'52606B',False)


def build():
 insets=make();cases={c:json.loads((V/c/'SUMMARY.json').read_text()) for c in ('C5','C7','L075')}
 author=json.loads((V/'author_L075/COMPARISON.json').read_text());latest=json.loads((V/'SUMMARY.json').read_text());core=cases['L075']
 assert latest['status']['ROOT_CAUSE_PRIMARY']=='INCONCLUSIVE'
 assert core['full_support_count']==27 and core['three_hop_closed_count']==0
 assert author['passed'] and author['force_relative_error']<1e-13
 prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9);slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=rgb('FFFFFF')
 prs.core_properties.title=C['title'].replace('\n',' ');prs.core_properties.subject='Current evidence and next experiment; no new numerical simulation';prs.core_properties.author='Clot CFD–PD / PFPD project'
 text(slide,C['title'],.43,.23,15.1,1.03,29,'203645',True)
 text(slide,C['subtitle'],.45,1.29,15.1,.35,15,'526879')
 line(slide,.45,1.72,15.55,1.72,'D4DDE4',.7)
 x=.40;top=1.90;gap=.15;height=2.63
 for stage in C['stages']:
  i=stage['id'];w=stage['width'];status='CURRENT' if i==C['current_stage'] else stage['status'];color=C['stage_colors'][status];fill=C['stage_fills'][status]
  shape(slide,MSO_SHAPE.RECTANGLE,x,top,w,height,fill,color,1.7 if status=='CURRENT' else .65,name=f'Stage {i}: {status}')
  shape(slide,MSO_SHAPE.RECTANGLE,x,top,w,.055,color,None,name=f'Stage {i} status rail')
  if C['show_status_icons']:
   shape(slide,MSO_SHAPE.OVAL,x+.12,top+.15,.23,.23,color,None,name=f'Stage {i} editable status icon')
   symbol='✓' if status=='COMPLETED' else '•' if status=='CURRENT' else '›' if status=='NEXT' else '○'
   text(slide,symbol,x+.12,top+.147,.23,.26,11,'FFFFFF',True,PP_ALIGN.CENTER)
  text(slide,f'{i}  {status}',x+.43,top+.17,w-.49,.24,11,color,True)
  text(slide,stage['label'],x+.12,top+.53,w-.24,1.06,16,color,True)
  body=C.get('stage_bodies_by_density',{}).get(C['visual_density'],{}).get(i,stage['body'][:2] if i==5 else stage['body'])
  body_size=11 if C['visual_density'] in ('LOW','HIGH') else 12
  text(slide,'\n'.join(body),x+.12,top+1.62,w-.24,.88,body_size,'344B5B')
  if C['show_numbers'].lower()!='none' and i==3:
   text(slide,f'Author force agreement:\n{author["force_relative_error"]*1e15:.1f} × 10⁻¹⁵ relative',x+.12,top+2.07,w-.24,.47,11.5,color)
  if C['show_numbers'].lower()!='none' and i==5:
   text(slide,f'L075: {core["full_support_count"]} full families;\nno closed force stencil',x+.12,top+2.07,w-.24,.47,11.5,color)
  if i==5:text(slide,'3δ: conservative bound',x,4.57,w,.25,10.9,'56636D')
  if i==6:text(slide,'Future; not yet validated',x+.05,4.57,w-.10,.25,11,'747E86')
  if i<6:shape(slide,MSO_SHAPE.CHEVRON,x+w+.022,top+1.08,.105,.23,'687580',None,name=f'Workflow arrow {i} to {i+1}')
  x+=w+gap
 text(slide,C['evidence_label'],.44,4.70,10.6,.24,10.5,'667B8A',True)
 slots=[(.40,4.05), (4.75,3.25), (8.30,3.75)]
 titles={'layer_forces':'(a) Affine is not equilibrium','direct_Z':'(b) Z-mediated force transfer','full_family':'(c) L = 0.75 mm: full family'}
 chosen=C['result_insets'][:C['max_result_insets']]
 for key,(xx,ww) in zip(chosen,slots):
  text(slide,titles[key],xx+.03,5.00,ww-.02,.29,12.5,'344B5B',True)
  slide.shapes.add_picture(str(insets[key]),Inches(xx),Inches(5.32),width=Inches(ww),height=Inches(2.10))
 if C['show_clot_schematic']:clot(slide,12.42,4.92,3.17)
 text(slide,C['caution_line'],.48,7.48,15.0,.25,11.5,'596D7A')
 # Three editable progress segments; the middle stage carries the focus.
 progx=[.45,5.57,10.70];progw=[4.85,4.85,4.85]
 for i,(xx,ww,label) in enumerate(zip(progx,progw,C['progress_labels'])):
  fill=['E8F1F7','FCE6CF','EEF0F2'][i];col=['0072B2','A6510D','75818B'][i]
  shape(slide,MSO_SHAPE.RECTANGLE,xx,7.91,ww,.44,fill,None,name='Progress: '+label)
  text(slide,label,xx+.06,7.995,ww-.12,.30,15,col,True,PP_ALIGN.CENTER)
  if i<2:shape(slide,MSO_SHAPE.CHEVRON,xx+ww+.08,8.035,.12,.19,'6D7B86',None,name='Progress arrow')
 text(slide,C['bottom_message'],.48,8.49,15.0,.29,14,'344B5B',True)
 text(slide,C['next_sentence'],.48,8.74,15.0,.23,11.5,'526879')
 # Notes are editable too and carry detail that would overcrowd a single slide.
 notes=make_notes(author,core,latest,chosen)
 slide.notes_slide.notes_text_frame.text=notes
 ppt=HERE/(NAME+'.pptx');prs.save(ppt)
 for rec in TEXT_RECORDS:
  assert rec['x']>=0 and rec['y']>=0 and rec['x']+rec['w']<=16.001 and rec['y']+rec['h']<=9.001
 reread=Presentation(ppt);assert len(reread.slides)==1
 (HERE/(NAME+'_notes.md')).write_text(notes)
 (HERE/'SLIDE_CONFIG.json').write_text(json.dumps(C,ensure_ascii=False,indent=2)+'\n')
 (HERE/'layout_validation.json').write_text(json.dumps({'slide_count':1,'aspect_ratio':'16:9','text_blocks':TEXT_RECORDS,'shape_count':len(slide.shapes),'raster_insets':len(chosen),'all_visible_text_English':True,'editable_workflow_text_shapes_arrows_and_clot_schematic':True,'whole_slide_rasterized':False,'scientific_simulations_run':0,'current_stage':C['current_stage'],'next_stage':5},ensure_ascii=False,indent=2)+'\n')
 print(ppt,flush=True)
 return ppt

def make_notes(author,core,latest,chosen):
 p=V/'REVIEW_BAQP_FINITE_DOMAIN_BOUNDARY_CONSISTENCY_V1.md';theory=V/'BOUNDARY_THEORY_NOTES.md'
 files=[str(p.relative_to(ROOT)),str(theory.relative_to(ROOT)),'../PROJECT_ACTIVE_CONTEXT.md','verification/pfpd_at1_c5_stability_v1/REVIEW_C5_STABILITY_V1.md','verification/pfpd_at1_c7_early_branch_v1/REVIEW_C7_EARLY_BRANCH_V1.md','verification/pfpd_at1_petsc_vi_solver_v1/REVIEW_PETSC_VI_SOLVER_V1.md',str((V/'SUMMARY.json').relative_to(ROOT)),str((V/'L075/SUMMARY.json').relative_to(ROOT)),str((V/'author_L075/COMPARISON.json').relative_to(ROOT))]
 inset=json.loads((HERE/'assets/inset_sources.json').read_text());files+=inset['source_files'];files+=inset['source_figure_counterparts']
 provenance=[]
 for rel in sorted(set(files)):
  f=ROOT/rel
  if not f.exists():raise FileNotFoundError(f)
  provenance.append(dict(path=rel,sha256=hashlib.sha256(f.read_bytes()).hexdigest(),bytes=f.stat().st_size))
 (HERE/'source_manifest.json').write_text(json.dumps(provenance,indent=2)+'\n')
 notes=['# Development Workflow and Current Progress','','## Objective','',C['primary_message'],'','One editable 16:9 slide, based only on archived repository evidence. No new mechanical solve, activation search, stability calculation or parameter sweep was run.','','## Source files','']+['- `'+r['path']+'`' for r in provenance]
 notes+=['','## Quantitative statements shown','',f'- Author implementation all-force relative difference: {author["force_relative_error"]:.16g}; shown as {author["force_relative_error"]*1e15:.1f} × 10⁻¹⁵.',f'- L075 denotes L={core["L_m"]*1000:g} mm: {core["full_support_count"]} self-family-complete points; {core["three_hop_closed_count"]} fully three-hop-closed points (also {core["two_hop_closed_count"]} two-hop-closed points).','- 3δ is the conservative geometric coefficient-dependency bound derived in the latest theory notes. It is a proposed isolation-test design, not an observed optimal physical boundary thickness.','- The three plots reproduce actual CSV/JSON coordinates and contributions. Their tick values are axis scales, not additional claimed experimental findings.','','## Stage interpretation','']
 for s in C['stages']:notes += [f'### {s["id"]}. {s["label"].replace(chr(10)," ")} — {s["status"]}','', '; '.join(s['body']+s['detail'])+'.','']
 notes+=['## Interpretation limits','','- Root cause remains formally INCONCLUSIVE. Boundary influence is a leading unresolved mechanism.','- No claim of a confirmed boundary artifact, wrong BAQP operator or incorrect surface BC is made.','- Geometric full family does not imply a closed force-dependency stencil.','- Agreement with the author code and energy gradient is an implementation check, not proof of physical traction correctness.','- C5 audited stable positive-phase states and C7 stationary states are different evidence; C7 stability is unassessed.','- Activation is not a crack prediction; stationarity is not stability.','- The clot drawing is a geometric schematic only: wall, attached mass, near-wall region and flow direction. No simulated contours, cracks or validated clot-wall failure are depicted.','- 3δ test is NEXT, not already run. Physical clot adhesion, detachment, rupture and loading calibration are FUTURE.','','## Insets and editability','','The same saved result panels are re-typeset from their original data, rather than cropping small raster labels that would be unreadable in projection. No interpolation, smoothing or fabricated trend is introduced. Workflow boxes, arrows, status icons, all narrative text, the progress indicator and the clot schematic are native editable PowerPoint objects; only the three real-data insets are images. PDF/SVG/PNG inset originals are included in assets/.','','## Steering and reproduction','','Edit `slide_config.py` (central `SLIDE_CONFIG`), then run:','','```bash','OPENBLAS_NUM_THREADS=1 OMP_NUM_THREADS=1 python generate_slide.py','```','','Controls include title, subtitle, current stage, stage labels/status/body, colors, density, inset selection (up to three), numeric callouts, schematic visibility and bottom messages. The default uses DejaVu Sans, the installed sans-serif font, to keep the editable slide and rendered preview consistent. Low density uses shorter body statements; high density uses more technical body statements. Both wordings are defined centrally in `stage_bodies_by_density`.','', 'The PNG is rendered from the actual PowerPoint via LibreOffice → PDF → PyMuPDF, not a separately reconstructed drawing. The temporary local LibreOffice runtime does not modify the system installation or scientific environment. If LibreOffice is available on PATH it is preferred; otherwise the current task runtime is used.','', 'The earlier 10-figure inventory and plan under `analysis/paper_figures/` were superseded by the user\'s one-slide steering before their figure generation. They remain as preparation, not as an additional claimed deliverable.']
 return '\n'.join(notes)+'\n'

def export(ppt):
 binary=shutil.which('libreoffice') or shutil.which('soffice');env=os.environ.copy();runtime=Path('/tmp/clot_slide_runtime_20261005/root')
 if not binary and (runtime/'usr/lib/libreoffice/program/soffice').exists():
  binary=str(runtime/'usr/lib/libreoffice/program/soffice');env['LD_LIBRARY_PATH']=str(runtime/'usr/lib/libreoffice/program')+':'+str(runtime/'usr/lib/x86_64-linux-gnu');env['SAL_USE_VCLPLUGIN']='svp'
 if not binary:raise RuntimeError('PowerPoint saved. LibreOffice is needed to render the actual slide.')
 command=[binary,'-env:UserInstallation=file:///tmp/clot_slide_lo_profile','--headless','--convert-to','pdf','--outdir',str(HERE),str(ppt)]
 proc=subprocess.run(command,env=env,capture_output=True,text=True,timeout=90);(HERE/'render.log').write_text(proc.stdout+proc.stderr)
 if proc.returncode:raise RuntimeError(proc.stdout+proc.stderr)
 pdf=HERE/(NAME+'.pdf');assert pdf.exists()
 try:import pymupdf
 except ImportError:
  # Reuse the existing local PDF reader without installing into the scientific environment.
  sys.path.insert(0,str(ROOT/'build_mechanical_equilibrium_v1/adr_completion/pdf_tools'));import pymupdf
 doc=pymupdf.open(pdf);assert len(doc)==1
 pix=doc[0].get_pixmap(dpi=300,alpha=False);pix.save(HERE/(NAME+'.png'))
 visible=doc[0].get_text();assert not re.search('[\u3400-\u9fff]',visible)
 (HERE/'rendered_slide_text.txt').write_text(visible)
 print('Rendered actual PowerPoint to PDF and PNG.',flush=True)
if __name__=='__main__':export(build())
