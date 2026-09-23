"""Show imposed boundary conditions alongside measured boundary flow rates."""
from pathlib import Path
import csv, hashlib, json, math, zipfile
import xml.etree.ElementTree as ET
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BASE=Path(__file__).resolve().parents[2]
CASE=BASE/'flow_cases/mean-2p0-mmps'
OUT=CASE/'presentation/pressure_flow_table'
STEM='boundary_pressure_flow'
HEADER='#5B9BD5'; STRIPES=['#D2DFF0','#E9EFF7']
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    source=CASE/'FLOW_2MMPS_VALIDATION.json'
    arrays=CASE/'frozen_flow/flow_arrays_si.npz'
    manifest=CASE/'frozen_flow/manifest.json'
    solver_xml=CASE/'run/solver.xml'
    sources={str(p.relative_to(CASE)):sha(p) for p in [source,arrays,manifest,solver_xml]}
    equation=ET.parse(solver_xml).find('.//Add_equation[@type="fluid"]')
    bcs={bc.attrib['name']:bc for bc in equation.findall('Add_BC')}
    data=json.loads(source.read_text());meta=json.loads(manifest.read_text())
    assert data['status']['FLOW_SOLVE']=='PASS' and data['target_inlet_mean_mm_s']==2
    assert sources['frozen_flow/flow_arrays_si.npz']==meta['files']['flow_arrays_si.npz']['sha256']
    m=data['measurements'];a=np.load(arrays)
    rows=[];display=[];mathdisplay=[];checks={}
    exponent_chars=str.maketrans('-0123456789','⁻⁰¹²³⁴⁵⁶⁷⁸⁹')
    for role,name in [('INLET','Inlet'),('OUTLET_01','Outlet 01'),('OUTLET_02','Outlet 02'),('OUTLET_03','Outlet 03')]:
        faces=a['boundary_triangles'][a['facet_tags']==meta['face_ids'][role]]
        points=a['points_m'][faces]
        area_vectors=.5*np.cross(points[:,1]-points[:,0],points[:,2]-points[:,0])
        area=np.linalg.norm(area_vectors,axis=1)
        signed=float(np.sum(area_vectors*a['velocity_m_s'][faces].mean(axis=1)))
        assert math.isclose(signed,m['signed_outward_boundary_flows_m3_s'][role],rel_tol=1e-12)
        bc=bcs[role];kind=bc.findtext('Type');value=float(bc.findtext('Value'))
        assert bc.findtext('Time_dependence')=='Steady'
        if role=='INLET':
            assert kind=='Dir' and bc.findtext('Impose_flux')=='true'
            assert math.isclose(-value,data['Q_target_m3_s'],rel_tol=1e-12)
            condition='Not prescribed\n(flow inlet)';traction=None;inlet_flow=-value
        else:
            assert kind=='Neu' and value==0
            condition='0 (zero traction)';traction=value;inlet_flow=None
        flow=abs(signed)*6e10;fraction=abs(signed)/m['Q_in_m3_s']*100
        mantissa,exponent=f'{flow:.3e}'.split('e');exponent=int(exponent)
        label=mantissa+' × 10'+str(exponent).translate(exponent_chars)
        pct='100%' if role=='INLET' else f'{fraction:.2f}%'
        values=[name,condition,label,pct]
        display.append(values)
        mathdisplay.append([name,values[1],rf'${mantissa}\times 10^{{{exponent}}}$',pct])
        rows.append(dict(boundary=role,solver_bc_type=kind,pressure_prescribed=False,
            normal_traction_bc_Pa=traction,prescribed_inlet_flow_m3_s=inlet_flow,signed_outward_flow_m3_s=signed,
            flow_magnitude_uL_min=flow,flow_relative_to_inlet_percent=fraction,
            displayed_pressure_traction_bc=condition,displayed_flow_uL_min=label,displayed_percentage=pct))
        checks[role]=dict(area_m2=float(area.sum()),boundary_condition_matches_solver_xml=True,flow_matches_saved_report=True)
    assert math.isclose(sum(r['flow_magnitude_uL_min'] for r in rows[1:]),rows[0]['flow_magnitude_uL_min'],rel_tol=1e-12)
    headers=['Boundary','Pressure / traction\nBC (Pa)','Volume flow\n(µL/min)','Flow relative to\ninlet']
    widths=[.18,.32,.25,.25]
    plt.rcParams.update({'font.family':'DejaVu Sans','mathtext.fontset':'dejavusans','pdf.fonttype':42,'svg.fonttype':'none'})
    fig,ax=plt.subplots(figsize=(12,3.55));fig.subplots_adjust(left=0,right=1,bottom=0,top=1);ax.axis('off')
    table=ax.table(cellText=mathdisplay,colLabels=headers,cellLoc='center',colLoc='center',
                   colWidths=widths,bbox=[.005,.015,.99,.97])
    table.auto_set_font_size(False)
    for (r,c),cell in table.get_celld().items():
        cell.set_edgecolor('white');cell.set_linewidth(1.5);cell.PAD=.02
        cell.set_height(.255 if r==0 else .18125)
        cell.set_facecolor(HEADER if r==0 else STRIPES[(r-1)%2])
        cell.get_text().set_fontsize(17 if r==0 or c==1 else 20)
        cell.get_text().set_color('white' if r==0 else 'black')
        cell.get_text().set_fontweight('bold' if r==0 else 'normal')
    fig.canvas.draw();renderer=fig.canvas.get_renderer()
    for cell in table.get_celld().values():
        cb=cell.get_window_extent(renderer);tb=cell.get_text().get_window_extent(renderer)
        assert cb.contains(tb.x0,tb.y0) and cb.contains(tb.x1,tb.y1)
    for ext in ['png','svg','pdf']:fig.savefig(OUT/f'{STEM}.{ext}',dpi=300,facecolor='white')
    plt.close(fig)
    with (OUT/f'{STEM}.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    for y,text,size,bold in [(.75,'Boundary conditions and volume flow',30,True),
                             (1.45,'Mean inlet speed: 2.0 mm/s | Imposed conditions and computed flow rates',18,False)]:
        box=slide.shapes.add_textbox(Inches(.65),Inches(y),Inches(12),Inches(.6))
        p=box.text_frame.paragraphs[0];p.text=text;p.font.size=Pt(size);p.font.bold=bold
    native=slide.shapes.add_table(5,4,Inches(.65),Inches(2.2),Inches(12.03),Inches(3.5)).table
    for column,width in zip(native.columns,widths):column.width=Inches(12.03*width)
    for r,values in enumerate([headers]+display):
        native.rows[r].height=Inches(.9 if r==0 else .65)
        for c,value in enumerate(values):
            cell=native.cell(r,c);cell.text=value;cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string((HEADER if r==0 else STRIPES[(r-1)%2])[1:])
            for p in cell.text_frame.paragraphs:
                p.alignment=PP_ALIGN.CENTER;p.font.name='Arial';p.font.size=Pt(19 if r==0 or c==1 else 23)
                p.font.bold=r==0;p.font.color.rgb=RGBColor.from_string('FFFFFF' if r==0 else '000000')
    box=slide.shapes.add_textbox(Inches(.65),Inches(5.98),Inches(12.03),Inches(.8))
    box.text_frame.text='BC = boundary condition. Inlet: prescribed flow; outlets: zero traction (Neumann).\nThe flow columns show computed magnitudes. Zero traction does not impose zero nodal pressure.'
    for p in box.text_frame.paragraphs:p.font.size=Pt(14);p.font.color.rgb=RGBColor.from_string('526172')
    slide.notes_slide.notes_text_frame.text='本表按用户要求显示施加的边界条件。入口采用 Dir + Impose_flux，指定流量，不指定压力；三个出口采用 Neu + Value=0，即零法向牵引边界（Pa）。出口的 0 是边界载荷设置，不是求解得到的截面平均压力，也不是压力节点强制为零。流量列及比例仍为当前 mean-2p0-mmps 工况 step 71 的计算结果；流量换算因子为 6 × 10^10 µL/min per m³/s。各出口比例分别四舍五入后合计 100.01%。Boundary conditions were read from the actual solver.xml; no solver parameters were changed.'
    ppt=OUT/f'{STEM}.pptx';prs.save(ppt)
    with zipfile.ZipFile(ppt) as z:assert z.testzip() is None
    recovered=next(s.table for s in Presentation(ppt).slides[0].shapes if s.has_table)
    assert [[recovered.cell(r,c).text for c in range(4)] for r in range(5)]==[headers]+display
    assert all(sha(CASE/p)==h for p,h in sources.items())
    result=dict(all_pass=True,case='mean-2p0-mmps',column_definition='IMPOSED_PRESSURE_OR_TRACTION_BOUNDARY_CONDITION',
                displays_solved_pressure=False,flow_columns='COMPUTED_FLOW_MAGNITUDES_AND_INLET_RELATIVE_PERCENTAGES',
                source_hashes=sources,independent_boundary_integration=checks,
                rows=rows,native_editable_pptx_table=True,pptx_text_roundtrip_pass=True,
                manual_powerpoint_UI_tested=False,table_text_fits_cells=True,
                output_sha256={f'{STEM}.{ext}':sha(OUT/f'{STEM}.{ext}') for ext in ['png','svg','pdf','csv','pptx']})
    (OUT/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(dict(rows=rows,output=str(OUT)),ensure_ascii=False,indent=2))


if __name__=='__main__':main()
