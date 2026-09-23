"""Assemble the checked figures, camera videos, PPT and Chinese review."""
from pathlib import Path
import hashlib
import json
import xml.etree.ElementTree as ET
import zipfile
from html.parser import HTMLParser
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from compute_field_diagnostics import ROOT, CASE, OUT, dump, sha


def main():
    compute=json.loads((OUT/'COMPUTE_VALIDATION.json').read_text())
    media=json.loads((OUT/'MEDIA_VALIDATION.json').read_text())
    visual=json.loads((OUT/'VISUAL_INSPECTION.json').read_text())
    nonlinear=json.loads((OUT/'NONLINEAR_RESIDUAL_AUDIT.json').read_text())
    residual_figures=json.loads((OUT/'RESIDUAL_FIGURE_VALIDATION.json').read_text())
    assert compute['all_pass'] and media['all_pass'] and visual['all_pass']
    assert nonlinear['all_pass'] and residual_figures['all_pass']
    for name,value in json.loads((OUT/'SOURCE_LOCK.json').read_text()).items():assert sha(CASE/name)==value
    for name,value in compute['outputs_sha256'].items():assert sha(OUT/name)==value
    for item in media['figures']+media['videos']:assert sha(OUT/item['file'])==item['sha256']
    tests=ET.parse(OUT/'tests.xml').getroot().findall('.//testcase')
    assert len(tests)>=42 and not any(c.find(tag) is not None for c in tests for tag in ['failure','error','skipped'])
    assert media['camera_position_constant'] and media['focal_point_constant'] and media['zoom_constant'] and media['axis_constant']
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    residual_notes=(
        '数据直接来自当前 2 mm/s 求解日志：71 个时间步、167 次非线性修正关联的线性求解。'
        '非线性主图 A 仅显示每步第一条残差的全程演化，B 显示 log10(本步第一条/本步最后一条) 的下降数量级。'
        'R_ref 是全程固定的第 1 步第 1 条残差；R 为求解器缩放系统的初始 KSP 范数，与原 NS 行定义一致。'
        '配对图每条竖线只连接同一个时间步内的首末记录，不把不同步的末值连成一条迭代历史。'
        '非线性停止：至少 2 次迭代，然后全程参考或步内参考任一相对残差达到 1e-10。'
        '第 4 步迭代数从前一步的 4 次变成 3 次，第 23 步从 3 次变成 2 次，均因较早满足停止条件；没有预条件器重建。'
        '低残差平台不代表 1e-14 的物理精度，也未被确认为浮点下限。'
        '诊断总图 C 为每次线性求解终止时的真实相对残差，D 为三次实际 GMRES 迭代历史。'
        '87 次按相对容差 1e-10 终止，80 次按绝对容差 1e-24 终止。'
        '后期相对残差高于 1e-10 不代表失败，全部真实残差满足相应综合停止阈值。'
        '线性残差对应 PETSc 内部缩放系统，不具有 Pa 单位。')
    for name in ['residual_nonlinear.png','residual_nonlinear_pairs.png','residual_convergence.png']:
        slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid()
        slide.background.fill.fore_color.rgb=RGBColor(255,255,255)
        path=OUT/'figures'/name
        with Image.open(path) as im:ratio=im.height/im.width
        width=int(min(Inches(12.9),Inches(7.1)/ratio));height=int(width*ratio)
        slide.shapes.add_picture(str(path),(prs.slide_width-width)//2,(prs.slide_height-height)//2,width=width,height=height)
        slide.notes_slide.notes_text_frame.text=residual_notes
    notes={
        'pressure':'显示原始血管表面的实际 FEM 压力（Pa），保持原压力参考；色标 -5 至 2500 Pa，未截去微小负值。',
        'wss':'WSS 由当前 P1 速度梯度、黏度与真实壁面法向后处理得到，未直接从求解器保存的 WSS 读取。'
              '只计算 WALL 面，不把出口和入口截面当血管壁。动画采用面面积加权的节点幅值显示；原始面值保存在 CSV/VTP。'
              '原始面最大值 51.142 Pa，节点显示最大值 46.516 Pa，二者应区分。未进行新增网格收敛分析。',
        'streamlines':'沿用已核验的 96 条真实 FEM 流线，出口 01/02/03 分别显示 16/56/24 条。条数不代表分流比例；未重新积分。'}
    for kind in ['pressure','wss','streamlines']:
        poster=OUT/'figures'/f'{kind}_overview_4k.png'
        for is_movie in [False,True]:
            slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid()
            slide.background.fill.fore_color.rgb=RGBColor(0,0,0)
            if is_movie:
                path=OUT/'animations'/f'{kind}_full_vessel_enlarged.mp4'
                slide.shapes.add_movie(str(path),0,0,prs.slide_width,prs.slide_height,
                    poster_frame_image=str(poster),mime_type='video/mp4')
            else:slide.shapes.add_picture(str(poster),0,0,width=prs.slide_width,height=prs.slide_height)
            slide.notes_slide.notes_text_frame.text=notes[kind]+(
                ' 相机、焦点和缩放全部固定。血管显示模型围绕穿过自身中心的固定长轴匀速转动；'
                '长轴仅从原始几何计算一次，原几何和场值不改变，整条血管不裁切。'
                '视频 1920×1080、24 fps、18 秒，是冻结稳态流场的展示旋转，不代表 18 秒瞬态求解。放映时点击视频播放。')
    deck=OUT/'Flow_Diagnostics_Presentation.pptx';prs.save(deck)
    expected={v['sha256'] for v in media['videos']}
    with zipfile.ZipFile(deck) as z:
        assert z.testzip() is None
        movies=[name for name in z.namelist() if name.startswith('ppt/media/') and name.endswith('.mp4')]
        assert len(movies)==3 and {hashlib.sha256(z.read(name)).hexdigest() for name in movies}==expected
    assert len(Presentation(deck).slides)==9
    c=compute;r=c['residuals'];w=c['wss_Pa']
    report=f'''# 完整血管流场：残差、压力与壁面剪切应力可视化

本次使用入口平均速度 **2.0 mm/s** 的已完成 FEM 流场，生成残差图、压力场及壁面剪切应力旋转动画，并同步放大原有流线动画。原求解结果、网格和原有流线均保持不变，没有重新运行 CFD 或积分微泡轨迹。

## 可直接使用的结果

- [统一浏览页面](OPEN_RESULTS.html)
- [9 页 PPT，内嵌三段视频](Flow_Diagnostics_Presentation.pptx)：非线性双面板主图、逐步首末配对图、残差诊断总图；压力静态图与视频；WSS 静态图与视频；流线静态图与视频。视频点击播放。
- [残差曲线 PNG](figures/residual_convergence.png)、[矢量 PDF](figures/residual_convergence.pdf)、[可编辑 SVG](figures/residual_convergence.svg)。单幅图另存为 `residual_nonlinear`、`residual_linear_termination`、`residual_gmres_histories` 的 PNG/PDF/SVG。
- [压力场旋转视频](animations/pressure_full_vessel_enlarged.mp4)、[4K 静态图](figures/pressure_overview_4k.png)。
- [壁面剪切应力旋转视频](animations/wss_full_vessel_enlarged.mp4)、[4K 静态图](figures/wss_overview_4k.png)。
- [放大的流线旋转视频](animations/streamlines_full_vessel_enlarged.mp4)、[4K 静态图](figures/streamlines_overview_4k.png)。

## 绘图与放大

残差图采用白色背景、英文标注、细线和简洁坐标轴，PNG 为 320 dpi，同时提供矢量 PDF/SVG。非线性主图拆成“每步初始残差的全程演化”和“本步下降数量级”两个面板，辅助配对图每条竖线仅对应同一步。全程参考量、实际非线性容差以及第 4、23 步的停止原因均已明确标注。低残差平台不解释为已证明的物理精度或浮点下限。详见 [残差专项说明与源码核验](RESIDUAL_REVIEW_ZH.md)。

三维图沿用此前流线图的黑底、turbo 配色、英文标题和右侧独立色标。本版取消随视角变化的自动居中与缩放。**相机位置、焦点、相机朝上方向和缩放比例全部固定**；只对显示对象应用绕固定轴的刚体旋转。旋转轴通过血管几何中心附近，并沿原始几何的主要延伸方向；轴在计算开始时确定一次，显示时保持竖直。三类场使用同一旋转轴、中心、角速度与相机。

轴的 SI 坐标系单位方向为 `{media['rotation_axis_unit']}`，固定中心为 `{media['rotation_center_um']}` µm，正交相机半高度固定为 **{media['fixed_parallel_scale_um']:.6f} µm**，角速度 **20°/s**。利用整个 360° 旋转的几何包络一次性确定尽可能大的固定尺度，完整血管约占三维视窗高度的 94.5%。没有加粗或修改血管几何。某些视角分支在投影中相互遮挡属于真实三维投影，旋转后可分辨。

上一个自动缩放版本及对应代码已备份到 `../field_diagnostics_revisions/before_fixed_axis/`，本次主目录、网页和 PPT 均指向新的固定轴版本。

每段视频 **1920×1080、24 fps、432 帧、18 秒**。这里的 18 秒是固定轴展示旋转的播放时间，不是瞬态血流演化时间。实际冻结结果为第 {c['last_step']} 步，时间步长 {c['dt_s']:.12e} s，结果时刻 {c['frozen_time_s']:.12e} s。

## 残差的含义与读取结果

数据来自 `../run/solver.log`，共 **{r['time_steps']} 个时间步、{r['linear_solves']} 次线性求解、{r['monitor_rows']:,} 行 PETSc 迭代监测记录**。重新读取原日志后，与已保存的求解历史逐条匹配。

非线性主图 A 只画每步第一条残差的全程演化，B 单独画 `log10(本步第一条/本步最后一条)`，以数量级表示步内下降量。辅助图用竖线配对同一步首末记录，避免把各步末值连成一条迭代历史。

全程固定参考为第 1 步第 1 条初始 KSP 范数，`R_ref = {nonlinear['reference_norm']:.12e}`。新图从更高精度的 KSP 记录重构比值，全部 167 条结果均与原 NS 行定义及打印值一致（允许原 NS 文本的四位有效数字舍入）。最后一条记录约为 **{r['final_nonlinear_Ri_over_R0']:.3e}**；其含义是最后一次修正关联的已记录残差，不是最终更新后额外评估的精确误差。

实际非线性停止条件是至少 2 次迭代后，全程参考或本步参考任一个比值达到 **1e-10**。图中的虚线表示全程参考这一项，不代表已确定的浮点下限。第 4、23 步分别以 3、2 次迭代满足条件，较前一步少一次；原始日志确认两处均无预条件器重建，线性求解均正常收敛。低残差不解释为 CFD 解的物理误差，网格与时间步精度仍需独立研究。

完整定义、源码位置、停止规则逐条验证及跳变证据见 [残差专项说明](RESIDUAL_REVIEW_ZH.md)、[审计 JSON](NONLINEAR_RESIDUAL_AUDIT.json)、[逐步 CSV](data/nonlinear_step_audit.csv)。

四面板总图 C 为每次线性求解结束时的 **真实相对残差**，D 为第 1 步第 1 次、第 30 步第 1 次、第 71 步第 2 次求解的实际 GMRES 迭代历史。87 次按相对容差 **1e-10** 终止，80 次按绝对容差 **1e-24** 终止。由于后期右端向量很小，达到绝对容差后，相对残差可以高于 1e-10；这不是求解失败。

所有线性求解的真实最终残差均小于 `max(rtol × 初始残差范数, atol)`；最大比值为 **{r['max_true_over_effective_tolerance']:.9f}**。日志没有失败线性求解、未识别 NS 行或恢复重试。线性残差属于 PETSc 内部缩放线性系统，不能解释为压力误差（Pa），残差收敛也不等于离散误差或网格收敛已经证明。

原始数据：[逐线性求解 CSV](data/residual_linear_solves.csv)、[全部真实残差迭代 CSV](data/residual_true_monitor.csv)、[逐时间步 CSV](data/residual_time_steps.csv)。

## 压力场

显示完整血管外表面及进出口截面上的原始 P1 节点压力；它是体压力场在边界上的取值。保留现有压力参考，未重新归零，也未把出口零牵引条件误解释成所有出口节点压力严格等于零。

原场压力范围 **{c['pressure_Pa']['min']:.6f}–{c['pressure_Pa']['max']:.6f} Pa**。色标固定为 **−5–2500 Pa**，所有帧一致，未裁去微小负压。压力数据：[SI 坐标 VTP](data/pressure_surface_si.vtp)。

## 壁面剪切应力

原冻结结果只有速度和压力，未直接保存 WSS。本次从原 **P1 四面体速度梯度**、**{c['viscosity_Pa_s']} Pa·s** 的动力黏度和真实 WALL 面法向计算切向黏性牵引的幅值。该定义与本项目 `svMultiPhysics` 的 `bpost` 一致；源码保存的向量采用相反符号约定，幅值不受影响。只处理 **{c['wall_facets']:,} 个 WALL 三角面**，入口和出口截面不作为血管壁。

每个壁面三角面都匹配到唯一相邻四面体；法向朝向流体域外侧。128 个单元还通过独立 4×4 仿射拟合复核，速度梯度最大绝对差为 **{c['independent_gradient_max_absolute_error_s_inv']:.3e} s⁻¹**。WSS 与法向的最大点积为 **{c['tangency_max_error_Pa']:.3e} Pa**。

| 量 | 数值（Pa） |
|---|---:|
| 原始三角面最小 WSS | {w['raw_min']:.6f} |
| 原始三角面最大 WSS | {w['raw_max']:.6f} |
| 壁面面积加权平均 WSS | {w['area_weighted_mean']:.6f} |
| 节点显示场最大 WSS | {w['display_max']:.6f} |

为避免 P1 单元梯度的片状颜色影响阅读，动画显示的是相邻三角面 **WSS 幅值按面积加权得到的节点场**，随后在面内插值。没有平滑几何，原始面值完整保存在 CSV/VTP；节点显示最大值与原始面最大值不同，论文统计应使用原始面值。固定色标 **0–55 Pa** 覆盖原始面值全范围。

这是已有解的 WSS 后处理，尚未新增网格收敛或独立 WSS 实验验证。当前结果是冻结稳态壁面剪切应力，不是脉动周期平均 WSS，也不包含 OSI。

原始数据：[逐壁面 WSS CSV](data/wall_wss_facets.csv)、[含原始面值、节点显示值及法向的 VTP](data/wall_wss_si.vtp)。

## 检查与复现

本项目 `tests/flow_2mmps` **{len(tests)} 项永久测试通过**。诊断部分包括 8 项 WSS 科学测试、3 项固定轴检查，本次增加 4 项残差语义测试，覆盖归一化不变性、两种参考的 OR 停止规则、最少迭代与上限区别以及第 4、23 步的日志核验。三段固定轴 MP4 沿用已逐帧解码验证的版本，视频哈希未改变。代理已查看新的残差图，科学解释的最终人工审核由用户完成。

PPTX 结构、9 页数量和 3 段内嵌视频的字节哈希通过检查；当前未进行原生 PowerPoint 界面试播。

数据计算、三维绘图、打包代码分别为 `scripts/flow_2mmps/compute_field_diagnostics.py`、`render_field_diagnostics.py`、`finalize_field_diagnostics.py`；仅重绘残差可运行 `scripts/flow_2mmps/residual_figures.py`。在项目根目录使用 `/home/lzy/projects/ulm_particle_3d_particle0/.venv/bin/python` 运行。

验证记录：[计算检查](COMPUTE_VALIDATION.json)、[媒体检查](MEDIA_VALIDATION.json)、[源文件哈希](SOURCE_LOCK.json)、[永久测试](tests.xml)、[视觉检查](VISUAL_INSPECTION.json)、[总验证](VALIDATION.json)。本次产物的 SHA-256 清单见 `SHA256SUMS.txt`。

定义依据：[本地 svMultiPhysics WSS 后处理源码](../../../vendor/svMultiPhysics_stage_q/Code/Source/solver/post.cpp)，[PETSc 真实残差监测](https://petsc.org/release/manualpages/KSP/KSPMonitorTrueResidual/)，[PETSc 线性求解与停止条件](https://petsc.org/release/manual/ksp/)。
'''
    (OUT/'REVIEW_ZH.md').write_text(report)
    sections=[]
    for kind,title,caption in [
        ('pressure','压力场','实际 FEM 压力，单位 Pa；完整血管表面显示。'),
        ('wss','壁面剪切应力','由当前速度梯度后处理获得，单位 Pa；原始面值和节点显示值均保留。'),
        ('streamlines','同步放大的流线','沿用原 96 条真实流线，三个出口均覆盖；未重新积分。')]:
        sections.append(f'''<section><h2>{title}</h2><p>{caption}</p>
<video controls loop playsinline preload="metadata" poster="figures/{kind}_overview_4k.png" src="animations/{kind}_full_vessel_enlarged.mp4"></video>
<p><a href="figures/{kind}_overview_4k.png">4K 静态图</a> · <a href="animations/{kind}_full_vessel_enlarged.mp4">下载 MP4</a> · <a href="figures/{kind}_rotation_views.png">四个旋转视角</a></p></section>''')
    page='''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>完整血管流场：残差、压力与壁面剪切应力</title><style>
body{margin:0;background:#10151e;color:#e5ebf4;font:17px/1.7 system-ui,sans-serif}main{max-width:1280px;margin:auto;padding:30px 24px}
h1{font-size:30px}h2{font-size:24px}a{color:#81c2ff}section{margin:30px 0;padding:24px;background:#19212d;border-radius:12px}
img,video{width:100%;height:auto;background:black;border-radius:4px}p{max-width:1000px}.notice{color:#bfcbdc}
</style><main><h1>完整血管流场：残差、压力与壁面剪切应力</h1>
<p>入口平均速度 2.0 mm/s 的已完成 FEM 流场。黑底三维可视化，整条血管放大显示；残差图采用白底学术样式。</p>
<p><a href="Flow_Diagnostics_Presentation.pptx">下载 PPT（9 页，内嵌三段视频）</a> · <a href="REVIEW_ZH.md">中文说明</a> · <a href="VALIDATION.json">验证记录</a></p>
<p class="notice">三段视频均绕穿过血管中心的固定长轴匀速旋转，相机、焦点和缩放保持不变。18 秒，24 fps；展示冻结稳态结果，不表示 18 秒瞬态求解。</p>
<section><h2>全程残差演化与单步内部下降</h2><img src="figures/residual_nonlinear.png" alt="非线性双面板图：全程初始残差演化及每步内部下降数量级">
<p>A：每步初始残差相对固定全程参考的变化。B：同一步首末残差下降的数量级。虚线是实际非线性容差；低残差平台不是已确认的浮点下限或物理精度。</p>
<p><a href="figures/residual_nonlinear.pdf">主图 PDF</a> · <a href="figures/residual_nonlinear.svg">主图 SVG</a> · <a href="RESIDUAL_REVIEW_ZH.md">源码与第 4、23 步核验说明</a></p></section>
<section><h2>逐步首末记录配对</h2><img src="figures/residual_nonlinear_pairs.png" alt="每个时间步的首末残差竖线配对图">
<p>每条竖线仅代表同一个时间步。第 4、23 步比前一步少一次迭代即满足停止条件，没有预条件器重建。</p>
<p><a href="figures/residual_nonlinear_pairs.pdf">配对图 PDF</a> · <a href="NONLINEAR_RESIDUAL_AUDIT.json">机器核验</a> · <a href="data/nonlinear_step_audit.csv">逐步下降量 CSV</a></p></section>
<section><h2>求解残差</h2><img src="figures/residual_convergence.png" alt="非线性残差、线性求解终止残差和 GMRES 历史">
<p><a href="figures/residual_convergence.pdf">矢量 PDF</a> · <a href="figures/residual_convergence.svg">SVG</a> · <a href="figures/residual_nonlinear.png">单独非线性残差图</a> · <a href="data/residual_true_monitor.csv">全部真实残差 CSV</a></p>
<p>71 个时间步，167 次线性求解。后期部分线性求解达到绝对容差而终止，图中的相对残差上升不表示求解失败。</p></section>
'''+''.join(sections)+'''<section><h2>原始计算数据</h2><p><a href="data/wall_wss_facets.csv">壁面 WSS CSV</a> · <a href="data/wall_wss_si.vtp">壁面 WSS VTP</a> · <a href="data/pressure_surface_si.vtp">压力表面 VTP</a> · <a href="COMPUTE_VALIDATION.json">计算验证</a> · <a href="MEDIA_VALIDATION.json">媒体验证</a></p></section></main></html>'''
    (OUT/'OPEN_RESULTS.html').write_text(page)
    validation=dict(all_pass=True,case='mean-2p0-mmps',tests_passed=len(tests),diagnostics_scientific_tests=8,fixed_axis_tests=3,new_residual_semantics_tests=4,
        residual_white_background=True,pressure_and_wss_match_streamline_visual_style=True,
        all_full_vessel_frames_unclipped=True,fixed_axis=True,constant_camera=True,constant_zoom=True,
        field_and_previous_streamlines_unchanged=True,source_lock_sha256=sha(OUT/'SOURCE_LOCK.json'),
        compute_validation_sha256=sha(OUT/'COMPUTE_VALIDATION.json'),media_validation_sha256=sha(OUT/'MEDIA_VALIDATION.json'),
        test_report_sha256=sha(OUT/'tests.xml'),visual_inspection_sha256=sha(OUT/'VISUAL_INSPECTION.json'),
        wss_is_derived_postprocessing=True,wss_mesh_convergence_not_assessed=True,
        nonlinear_reference_fixed=True,nonlinear_stop_rule_source_audited=True,numerical_floor_not_claimed=True,
        nonlinear_audit_sha256=sha(OUT/'NONLINEAR_RESIDUAL_AUDIT.json'),residual_figures_validation_sha256=sha(OUT/'RESIDUAL_FIGURE_VALIDATION.json'),
        videos=media['videos'],presentation=dict(file=deck.name,slides=9,embedded_videos=3,sha256=sha(deck),
            native_powerpoint_ui_tested=False),manual_scientific_review='PENDING_USER_REVIEW',script_sha256=sha(__file__))
    dump(OUT/'VALIDATION.json',validation)
    class Links(HTMLParser):
        def handle_starttag(self,tag,attrs):
            for key,value in attrs:
                if key in ['href','src','poster']:assert (OUT/value).is_file(), value
    Links().feed(page)
    files=sorted(p for p in OUT.rglob('*') if p.is_file() and p.name!='SHA256SUMS.txt')
    (OUT/'SHA256SUMS.txt').write_text(''.join(f'{sha(p)}  {p.relative_to(OUT)}\n' for p in files))
    print(json.dumps(dict(all_pass=True,tests=len(tests),files=len(files),presentation=str(deck)),indent=2))


if __name__=='__main__':main()
