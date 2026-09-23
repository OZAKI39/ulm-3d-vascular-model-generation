#!/usr/bin/env python3
"""Export the measured 2 mm/s case boundary flows in microlitres per minute."""
from pathlib import Path
import csv,hashlib,json,math,zipfile
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from PIL import Image
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR

BASE=Path(__file__).resolve().parents[2]
CASE=BASE/'flow_cases/mean-2p0-mmps'
OUT=CASE/'presentation';STEM='boundary_flow_uL_min'
CONVERSION=60*10**9  # 1 m^3 = 10^9 microlitres; 1 minute = 60 seconds.
HEADER='#5B9BD5';STRIPES=['#D2DFF0','#E9EFF7']


def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def main():
    OUT.mkdir(exist_ok=True)
    source=CASE/'FLOW_2MMPS_VALIDATION.json';csvsource=CASE/'reports/boundary_flows.csv'
    before={str(p):sha(p) for p in [source,csvsource]}
    data=json.loads(source.read_text());m=data['measurements']
    assert data['case']=='mean-2p0-mmps' and data['status']['FLOW_SOLVE']=='PASS'
    assert data['target_inlet_mean_mm_s']==2.0
    with csvsource.open(newline='') as f:independent={r['boundary']:float(r['signed_outward_m3_s']) for r in csv.DictReader(f)}
    order=['INLET','OUTLET_01','OUTLET_02','OUTLET_03'];names=['Inlet','Outlet 01','Outlet 02','Outlet 03']
    rows=[];display=[];mathdisplay=[]
    superscript=str.maketrans('-0123456789','⁻⁰¹²³⁴⁵⁶⁷⁸⁹')
    for boundary,name in zip(order,names):
        signed=m['signed_outward_boundary_flows_m3_s'][boundary]
        assert math.isclose(signed,independent[boundary],rel_tol=1e-14,abs_tol=0)
        q=abs(signed);converted=q*CONVERSION;fraction=100*q/m['Q_in_m3_s']
        mantissa,exponent=f'{converted:.3e}'.split('e');exponent=int(exponent)
        label=mantissa+' × 10'+str(exponent).translate(superscript)
        percentage='100%' if boundary=='INLET' else f'{fraction:.2f}%'
        rows.append(dict(boundary=boundary,signed_outward_m3_s=signed,flow_magnitude_m3_s=q,
            flow_magnitude_uL_min=converted,flow_relative_to_inlet_percent=fraction,
            displayed_flow_uL_min=label,displayed_percentage=percentage))
        display.append([name,label,percentage])
        mathdisplay.append([name,rf'${mantissa}\times 10^{{{exponent}}}$',percentage])
    assert math.isclose(sum(r['flow_magnitude_uL_min'] for r in rows[1:]),rows[0]['flow_magnitude_uL_min'],rel_tol=1e-12)
    with (OUT/f'{STEM}.csv').open('w',newline='',encoding='utf-8-sig') as f:
        writer=csv.DictWriter(f,fieldnames=list(rows[0]));writer.writeheader();writer.writerows(rows)
    headers=['Boundary','Volume flow (µL/min)','Flow relative to inlet']
    plt.rcParams.update({'font.family':'DejaVu Sans','mathtext.fontset':'dejavusans','pdf.fonttype':42,'svg.fonttype':'none'})
    fig,ax=plt.subplots(figsize=(12,3.3));fig.subplots_adjust(left=0,right=1,bottom=0,top=1);ax.axis('off')
    table=ax.table(cellText=mathdisplay,colLabels=headers,cellLoc='center',colLoc='center',
        colWidths=[.265,.38,.355],bbox=[.005,.015,.99,.97])
    table.auto_set_font_size(False);table.set_fontsize(19)
    for (r,c),cell in table.get_celld().items():
        cell.set_edgecolor('white');cell.set_linewidth(2);cell.PAD=.03
        cell.set_facecolor(HEADER if r==0 else STRIPES[(r-1)%2])
        if r==0:cell.get_text().set_color('white');cell.get_text().set_fontweight('bold');cell.get_text().set_fontsize(18)
    fig.canvas.draw();renderer=fig.canvas.get_renderer()
    for cell in table.get_celld().values():
        cb=cell.get_window_extent(renderer);tb=cell.get_text().get_window_extent(renderer)
        assert cb.contains(tb.x0,tb.y0) and cb.contains(tb.x1,tb.y1)
    for fmt in ['png','svg','pdf']:fig.savefig(OUT/f'{STEM}.{fmt}',dpi=400,facecolor='white')
    plt.close(fig)
    # A native, editable table for direct copy/paste into the user's presentation.
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    title=slide.shapes.add_textbox(Inches(.7),Inches(.85),Inches(12),Inches(.65))
    p=title.text_frame.paragraphs[0];p.text='Boundary volume flow';p.font.size=Pt(30);p.font.bold=True
    sub=slide.shapes.add_textbox(Inches(.7),Inches(1.5),Inches(12),Inches(.5))
    p=sub.text_frame.paragraphs[0];p.text='New flow field | Mean inlet speed: 2.0 mm/s';p.font.size=Pt(18)
    p.font.color.rgb=RGBColor.from_string('526172')
    shape=slide.shapes.add_table(5,3,Inches(.7),Inches(2.25),Inches(11.93),Inches(3.2))
    native=shape.table
    for col,width in zip(native.columns,[.265,.38,.355]):col.width=Inches(11.93*width)
    for r,values in enumerate([headers]+display):
        for c,value in enumerate(values):
            cell=native.cell(r,c);cell.text=value;cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string((HEADER if r==0 else STRIPES[(r-1)%2])[1:])
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment=PP_ALIGN.CENTER;paragraph.font.name='Arial';paragraph.font.size=Pt(22)
                paragraph.font.bold=r==0;paragraph.font.color.rgb=RGBColor.from_string('FFFFFF' if r==0 else '111111')
    foot=slide.shapes.add_textbox(Inches(.7),Inches(5.7),Inches(12),Inches(.65))
    p=foot.text_frame.paragraphs[0];p.text='Flow magnitudes are shown; each percentage is relative to the measured inlet flow.'
    p.font.size=Pt(14);p.font.color.rgb=RGBColor.from_string('526172')
    slide.notes_slide.notes_text_frame.text='Source: mean-2p0-mmps/FLOW_2MMPS_VALIDATION.json, actual VTU boundary surface integration. Conversion: Q [µL/min] = Q [m³/s] × 6 × 10^10. The signed outward inlet flux is negative; the table shows its magnitude. Independently rounded percentages sum to 100.01%.'
    ppt=OUT/f'{STEM}.pptx';prs.save(ppt)
    with zipfile.ZipFile(ppt) as z:assert z.testzip() is None
    assert len(Presentation(ppt).slides)==1
    assert all(sha(p)==h for p,h in before.items())
    with Image.open(OUT/f'{STEM}.png') as im:dimensions=list(im.size);im.verify()
    result=dict(all_pass=True,source_case=data['case'],source_sha256=before,source_files_unchanged=True,
        target_inlet_mean_mm_s=2.0,conversion_factor_m3_s_to_uL_min=CONVERSION,
        source_quantities='ACTUAL_VTU_BOUNDARY_SURFACE_INTEGRALS',display_uses_flow_magnitudes=True,
        percentage_denominator='ACTUAL_INLET_FLOW_MAGNITUDE',rows=rows,png_dimensions=dimensions,
        table_text_fits_cells=True,pptx_native_editable_table=True,manual_powerpoint_UI_tested=False,
        output_sha256={f'{STEM}.{fmt}':sha(OUT/f'{STEM}.{fmt}') for fmt in ['png','pdf','svg','csv','pptx']})
    (OUT/f'{STEM}_validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(dict(rows=rows,outputs=str(OUT)),ensure_ascii=False,indent=2))


if __name__=='__main__':main()
