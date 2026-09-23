"""Add the verified streamline section to the existing case and make PPT assets."""
from pathlib import Path
import json,shutil,zipfile,xml.etree.ElementTree as ET
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from compute_streamlines import CASE,OUT,ROOT,sha,dump


def main():
    compute=json.loads((OUT/'COMPUTE_VALIDATION.json').read_text())
    media=json.loads((OUT/'MEDIA_VALIDATION.json').read_text())
    visual=json.loads((OUT/'VISUAL_INSPECTION.json').read_text())
    assert compute['all_pass'] and media['all_pass'] and visual['all_pass']
    tests=ET.parse(OUT/'final_tests.xml').getroot().findall('.//testcase')
    assert len(tests)>=5 and not any(c.find('failure') is not None or c.find('error') is not None or c.find('skipped') is not None for c in tests)
    for file,value in json.loads((OUT/'SOURCE_LOCK.json').read_text()).items():assert sha(CASE/file)==value
    for record in media['figures']+[media['video']]:assert sha(CASE/record['file'])==record['sha256']
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    for name in ['Figure_06_streamlines_overview.png','MOVIE','Figure_07_streamlines_outlet_coverage.png']:
        slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor(0,0,0)
        if name=='MOVIE':
            slide.shapes.add_movie(str(CASE/media['video']['file']),0,0,prs.slide_width,prs.slide_height,
                poster_frame_image=str(OUT/'inspection/frame_0144.png'),mime_type='video/mp4')
        else:slide.shapes.add_picture(str(CASE/'figures'/name),0,0,width=prs.slide_width,height=prs.slide_height)
        slide.notes_slide.notes_text_frame.text='入口均速 2.0 mm/s 的新 FEM 流场。96 条实际积分流线，出口 01/02/03 分别选取 16/56/24 条以保证视觉覆盖；条数不是分流比例。颜色表示保存流场沿流线的速度，单位 mm/s。这是点流线，不是有限尺寸微泡轨迹。视频为稳态流场的 18 秒相机环绕，放映时点击画面播放。'
    deck=OUT/'Streamlines_2mmps_Presentation.pptx';prs.save(deck)
    with zipfile.ZipFile(deck) as z:
        assert z.testzip() is None
        movies=[n for n in z.namelist() if n.startswith('ppt/media/') and n.endswith('.mp4')]
        assert len(movies)==1
        import hashlib
        assert hashlib.sha256(z.read(movies[0])).hexdigest()==media['video']['sha256']
    assert len(Presentation(deck).slides)==3
    counts=compute['candidate_outcomes']
    review=f'''# 新流场流线可视化：中文审核说明

本次在入口平均速度 **2.0 mm/s** 的新 FEM 流场上补充流线显示。保持黑底、半透明灰色血管、按速度着色的内部流线和右侧独立色标，色标固定为 **0–7.5 mm/s**。原流场、网格、既有静态图和两段视频未改变。

## 展示内容与出口覆盖

从入口按实际正向通量密度生成 768 个候选点，沿保存的 P1 速度场正向积分。自然到达出口 01、02、03 的候选数分别为 **{counts['OUTLET_01']}、{counts['OUTLET_02']}、{counts['OUTLET_03']}**；另有 **{counts.get('POINT_TIME_HORIZON_30S',0)}** 条达到候选流线的 30 秒积分上限，未作为已连通流线显示。

最终展示 **96 条**：出口 01 为 **16 条**，出口 02 为 **56 条**，出口 03 为 **24 条**。在各出口已完成的真实流线中，通过末端位置的最远点选取增加截面覆盖，降低重复线束堆积。没有指定流线必须流向某个出口，也没有修改速度、补接路径或将曲线平滑后代替积分结果。

这是为兼顾小流量分支可读性的分层展示，**显示条数不代表分流比例**。实际新场出口比例仍约为 **4.26%、85.21%、10.54%**。入口位置最初按通量采样，但显示集合经过出口分层与空间间距筛选，不能当作无偏通量样本。

## 直接使用的文件

- [完整血管流线图，4K](../figures/Figure_06_streamlines_overview.png)：标出入口和三个出口。
- [各出口流线覆盖图，4K](../figures/Figure_07_streamlines_outlet_coverage.png)：总览与三个出口分别显示。
- [18 秒旋转动画](../animations/Animation_03_streamlines_full_vessel.mp4)：1920×1080，24 fps，432 帧。
- [视频关键帧总览](../figures/Figure_08_streamlines_storyboard.png)：直接解码最终视频生成。
- [3 页 PPT，视频已内嵌](Streamlines_2mmps_Presentation.pptx)：可插入原汇报，视频点击播放。
- [流线 CSV](data/catalog.csv)、[SI 单位 VTP](data/streamlines_si.vtp)、[原始与细化路径](data/selected_paths.npz)。

## 科学含义与计算方式

这些是冻结稳态速度场的**点流线**。它们描述血流方向及连通关系，未加入微泡半径、近壁颗粒阻力或接纳筛选，因此不同于前面的有限尺寸微泡轨迹。

使用现有双精度 P1 RK23 点示踪器、四面体邻接搜索和原始官方出口三角面。最大空间步长 0.15 µm，局部位置误差控制值 0.00001 µm。末端使用沿局部速度的小段积分与真实出口面的交点，末段长度上限为 3 nm；这是出口事件的数值近似，不是凭目标出口连线。所有保存路径点均在原四面体网格内，所有选中流线的末段均重新检验其出口三角面。

每条选中流线另外以空间步长减半、局部误差控制值缩小至四分之一重新积分，出口归属均一致。按归一化弧长对齐后的最大路径差异为 **{compute['maximum_refined_path_difference_um']:.6f} µm**，最大出口位置差异为 **{compute['maximum_refined_endpoint_difference_um']:.6f} µm**。该检查支持当前可视化的几何稳定性，不替代独立网格收敛研究。

速度颜色取自实际 P1 场插值。永久测试使用独立的重心坐标线性求解核验保存速度，并严格判断四面体包含关系，避免通用点探针在单元边界附近的宽松容差选择邻单元。

模型和流线的物理坐标保持不变，显示时只作一致的米到微米转换。线宽服务于显示，没有粒径含义。动画只旋转相机，**18 秒是视频时长，不是一次非稳态血流演化的时长**。

## 验证结果与复现

本轮计算在 WSL 主机 `{compute['hostname']}` 完成，小规模候选积分、筛选及细化复核共约 {compute['elapsed_seconds']:.2f} 秒。渲染使用本机 Mesa / llvmpipe，耗时约 {media['render_seconds']:.1f} 秒，未启动新的 CFD 求解或大规模微泡积分。

当前 `tests/flow_2mmps` 测试集 **{len(tests)} 项通过**，包括 5 项新增流线检查。实际 MP4 的 432 帧全部解码，432 帧均不同；逐帧检查完整血管位于视窗内。代理已查看总览、各出口分图和视频关键帧。原 {compute['unchanged_file_count']} 个流场、网格及旧可视化文件哈希不变。

PPTX 结构和内嵌视频字节通过检查；当前环境没有原生 PowerPoint，未宣称完成 PowerPoint 界面试播。用户科学审核保持待确认。

新增代码：`scripts/flow_2mmps/compute_streamlines.py`、`render_streamlines.py`、`finalize_streamlines.py`；新增永久测试：`tests/flow_2mmps/test_streamlines.py`。复用的点示踪代码位于 `/home/lzy/projects/ulm_particle_3d_particle0/particle_3d/src`，依赖文件哈希全部保存在 [COMPUTE_VALIDATION.json](COMPUTE_VALIDATION.json)。完成积分后脚本拒绝直接重算覆盖；再渲染可直接运行 `render_streamlines.py`。

`AUTOMATED_CHECKS = PASS`；`ALL_OUTLETS_COVERED = PASS`；`ANIMATION_DECODE = PASS`；`MANUAL_VISUAL_REVIEW = PENDING_USER_REVIEW`。
'''
    (OUT/'STREAMLINES_REVIEW.md').write_text(review)
    validation=dict(all_pass=True,field_case='mean-2p0-mmps',selected_streamlines=96,
        outlet_counts=compute['selected_outlet_counts'],line_counts_not_flow_fractions=True,
        tests_passed=len(tests),test_report_sha256=sha(OUT/'final_tests.xml'),
        computation_record_sha256=sha(OUT/'COMPUTE_VALIDATION.json'),media_record_sha256=sha(OUT/'MEDIA_VALIDATION.json'),
        visual_inspection=visual,source_field_sha256=compute['source_field_sha256'],
        video=media['video'],figures=media['figures'],
        pptx=dict(path=deck.name,slides=3,embedded_videos=1,sha256=sha(deck),native_powerpoint_UI_tested=False),
        manual_scientific_review='PENDING_USER_REVIEW')
    dump(OUT/'STREAMLINES_VALIDATION.json',validation)
    # Preserve the previous index/report/checksum manifest before adding a section.
    for name in ['OPEN_RESULTS.html','FLOW_2MMPS_REVIEW.md','OUTPUT_SHA256SUMS.txt']:
        backup=OUT/('before_streamlines_'+name)
        if not backup.exists():shutil.copy2(CASE/name,backup)
    page=(OUT/'before_streamlines_OPEN_RESULTS.html').read_text()
    section='''<section id="streamlines"><h2>新增：流线与全部出口覆盖</h2>
<p>新流场入口均速 2.0 mm/s。96 条实际积分流线覆盖全部出口：出口 01 / 02 / 03 分别显示 16 / 56 / 24 条。为清楚展示小流量分支，流线按出口分层选取，条数不代表分流比例。</p>
<p><a href="streamlines/Streamlines_2mmps_Presentation.pptx">下载流线 PPT（3 页，内嵌视频）</a> · <a href="streamlines/STREAMLINES_REVIEW.md">流线说明</a> · <a href="streamlines/STREAMLINES_VALIDATION.json">验证记录</a></p>
<img loading="lazy" src="figures/Figure_06_streamlines_overview.png" alt="完整血管流线与三个出口">
<video controls loop preload="metadata" src="animations/Animation_03_streamlines_full_vessel.mp4"></video>
<img loading="lazy" src="figures/Figure_07_streamlines_outlet_coverage.png" alt="各出口流线覆盖">
<img loading="lazy" src="figures/Figure_08_streamlines_storyboard.png" alt="流线旋转动画关键帧">
</section><h2>原有流场可视化</h2>'''
    page=page.replace('<video',section+'<video',1);(CASE/'OPEN_RESULTS.html').write_text(page)
    original=(OUT/'before_streamlines_FLOW_2MMPS_REVIEW.md').read_text()
    (CASE/'FLOW_2MMPS_REVIEW.md').write_text(original+'\n\n## 补充：流线可视化与全部出口覆盖\n\n新增 96 条新流场流线、4K 总览和各出口分图、18 秒旋转动画，以及可直接使用的 PPT。三个出口分别显示 16 / 56 / 24 条；显示数量用于保证可读性，不代表分流比例。见 [流线专项说明](streamlines/STREAMLINES_REVIEW.md) 与 [验证记录](streamlines/STREAMLINES_VALIDATION.json)。\n')
    entries={line.split('  ',1)[1]:line.split('  ',1)[0] for line in (OUT/'before_streamlines_OUTPUT_SHA256SUMS.txt').read_text().splitlines()}
    added=[CASE/r['file'] for r in media['figures']+[media['video']]]+[p for p in OUT.rglob('*') if p.is_file()]
    for p in added+[CASE/'OPEN_RESULTS.html',CASE/'FLOW_2MMPS_REVIEW.md']:entries[str(p.relative_to(CASE))]=sha(p)
    (CASE/'OUTPUT_SHA256SUMS.txt').write_text(''.join(f'{value}  {name}\n' for name,value in sorted(entries.items())))
    assert all(sha(CASE/name)==value for name,value in entries.items())
    print(json.dumps(dict(all_pass=True,tests=len(tests),checksum_files=len(entries),presentation=str(deck)),indent=2))


if __name__=='__main__':main()
