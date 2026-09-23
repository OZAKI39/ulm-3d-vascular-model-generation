"""Audit and package common-clock translation and rotation for a PPT."""
import json
import hashlib
import zipfile
import xml.etree.ElementTree as ET
import numpy as np
from scipy.linalg import expm
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from prepare_rbc_mb_flow_rotation import OUT, ROOT, sha, write_json


def rotation_distance(a,b):
    minus=np.linalg.norm(a-b,axis=-1);plus=np.linalg.norm(a+b,axis=-1)
    return 4*np.arctan2(np.minimum(minus,plus),np.maximum(minus,plus))


def audit_motion(scene,data):
    q=data['quaternion_wxyz'];t=data['time_s'];g=np.array(scene['records'][0]['gradient_s_inv'])
    lam=scene['rbc_jeffery_lambda'];e=.5*(g+g.T);w=.5*(g-g.T)
    exact=[]
    for time in t:
        p=(expm((w+lam*e)*time)@data['rbc_short_axis'][0].T).T
        exact.append(p/np.linalg.norm(p,axis=1)[:,None])
    exact=np.asarray(exact);p=data['rbc_short_axis']
    axis_errors=np.arctan2(np.linalg.norm(np.cross(p,exact),axis=-1),np.sum(p*exact,axis=-1))
    change=np.arctan2(np.linalg.norm(np.cross(p,p[0]),axis=-1),np.abs(np.sum(p*p[0],axis=-1)))
    omega=np.linalg.norm(data['angular_velocity_s_inv'],axis=-1)
    radii=np.array([r['radius_m'] for r in scene['records']]);initial=data['positions_m'][0]
    wall=np.min(scene['tube_radius_m']-np.linalg.norm(initial[:,1:],axis=1)-radii)
    gaps=[];vel=data['velocity_m_s']
    for j in range(len(radii)):
        for k in range(j):
            d=initial[j]-initial[k];v=vel[j]-vel[k]
            tm=0 if v@v==0 else np.clip(-d@v/(v@v),0,t[-1])
            gaps.append(np.linalg.norm(d+tm*v)-radii[j]-radii[k])
    result=dict(
        maximum_refinement_orientation_error_deg=float(np.rad2deg(rotation_distance(q,data['quaternion_coarse_wxyz'])).max()),
        maximum_exact_Jeffery_axis_error_deg=float(np.rad2deg(axis_errors).max()),
        maximum_quaternion_norm_error=float(np.abs(np.linalg.norm(q,axis=-1)-1).max()),
        each_RBC_maximum_axis_change_deg=np.rad2deg(change).max(axis=0).tolist(),
        RBC_angular_speed_range_rad_s=[float(omega[:,:9].min()),float(omega[:,:9].max())],
        MB_angular_speed_range_rad_s=[float(omega[:,9:].min()),float(omega[:,9:].max())],
        MB_accumulated_rotation_deg=(np.rad2deg(omega[-1,9:]*t[-1])).tolist(),
        minimum_all_orientation_wall_bound_um=float(wall*1e6),
        minimum_all_time_bounding_sphere_pair_gap_um=float(min(gaps)*1e6),
        original_shape_sha256_verified=sha(OUT/'data/normal_rbc_si.vtp')==scene['shape_source_sha256'],
    )
    assert result['maximum_refinement_orientation_error_deg']<.12
    assert result['maximum_exact_Jeffery_axis_error_deg']<.18
    assert result['maximum_quaternion_norm_error']<5e-15
    assert min(result['each_RBC_maximum_axis_change_deg'])>30
    assert wall>0 and min(gaps)>0 and result['original_shape_sha256_verified']
    write_json(OUT/'data/MOTION_VERIFICATION.json',result)
    # Compact numerical reference figure, useful for scientific review.
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    with plt.rc_context({'font.family':'DejaVu Sans','font.size':11,'axes.spines.top':False,'axes.spines.right':False}):
        fig,ax=plt.subplots(1,2,figsize=(10.8,4),constrained_layout=True)
        for j,c in enumerate(['#c64545','#d0802c','#7e64a8']):
            ax[0].plot(t*1000,np.rad2deg(change[:,j]),c=c,label=f'RBC initial orientation {j+1}')
        ax[0].set(xlabel='Physical time (ms)',ylabel='Short-axis change (degrees)',title='A  RBC reorientation')
        ax[0].legend(fontsize=8,frameon=False);ax[0].set_ylim(0,95)
        for j,c in [(9,'#078fa9'),(23,'#193c69')]:
            ax[1].plot(t*1000,np.rad2deg(omega[:,j]*t),c=c,ls='-' if j==9 else '--',label=f'MB lane {1 if j==9 else 2}')
        ax[1].set(xlabel='Physical time (ms)',ylabel='Accumulated rotation (degrees)',title='B  MB rotation from local vorticity')
        ax[1].legend(frameon=False)
        for a in ax:a.grid(alpha=.18)
        fig.savefig(OUT/'figures/rotation_verification.png',dpi=250)
        fig.savefig(OUT/'figures/rotation_verification.svg');plt.close(fig)
    return result


def main():
    read=lambda name:json.loads((OUT/name).read_text())
    scene=read('data/SCENE.json');media=read('MEDIA_VALIDATION.json');inspection=read('VISUAL_INSPECTION.json')
    data=np.load(OUT/'data/motion.npz');motion=audit_motion(scene,data)
    suites=list(ET.parse(OUT/'tests.xml').getroot().iter('testsuite'))
    count=sum(int(s.get('tests',0)) for s in suites)
    assert count==13 and all(int(s.get(k,0))==0 for s in suites for k in ['failures','errors','skipped'])
    assert media['all_pass'] and inspection['all_pass']
    assert sha(OUT/'data/motion.npz')==scene['array_sha256']==media['source_npz_sha256']
    for name,digest in inspection['inspected_image_sha256'].items():assert sha(OUT/name)==digest
    video=OUT/'animations/RBC_MB_flow_rotation.mp4'
    assert sha(video)==media['video_sha256']
    deck=OUT/'RBC_MB_Flow_Rotation_Presentation.pptx'
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    notes=('RBC 与 MB 在同一解析 Poiseuille 流场中的平移与旋转。血管直径 14 µm，平均流速 2 mm/s；'
           '本理想化场具备明确的速度、梯度和涡量，已导出 VTU。正常双凹红细胞保持形状，'
           '旋转沿用体积匹配扁球体的 Jeffery 近似；微泡沿用球形粒子半涡量角速度。'
           '微泡表面的白色条带和金色标记仅用于显示已计算的姿态，不表示真实膜纹理。'
           '物理时间 30 ms，播放 18 秒；9 个 RBC 和 28 个 MB 是展示队列，不定义浓度。'
           '这是给定背景流驱动的独立粒子运动，不包含粒子间流体反馈或膜变形求解。')
    for movie in [False,True]:
        slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid()
        slide.background.fill.fore_color.rgb=RGBColor(0,0,0)
        if movie:
            slide.shapes.add_movie(str(video),0,0,prs.slide_width,prs.slide_height,
                poster_frame_image=str(OUT/'figures/RBC_MB_flow_rotation_overview.png'),mime_type='video/mp4')
        else:
            slide.shapes.add_picture(str(OUT/'figures/RBC_MB_flow_rotation_overview.png'),0,0,
                width=prs.slide_width,height=prs.slide_height)
        slide.notes_slide.notes_text_frame.text=notes+(' 放映时点击视频播放。' if movie else '')
    prs.save(deck)
    with zipfile.ZipFile(deck) as z:
        assert z.testzip() is None
        videos=[n for n in z.namelist() if n.endswith('.mp4')]
        assert len(videos)==1 and hashlib.sha256(z.read(videos[0])).hexdigest()==media['video_sha256']
    assert len(Presentation(deck).slides)==2
    sources=['particle_3d/src/particle_3d/'+n for n in ['coflow_rotation.py','rbc_orientation.py','rbc_integrator.py','microbubble.py','normal_rbc_encounter.py']]
    sources+=['particle_3d/scripts/'+n for n in ['prepare_rbc_mb_flow_rotation.py','render_rbc_mb_flow_rotation.py','finalize_rbc_mb_flow_rotation.py']]
    sources+=['particle_3d/tests/rbc_mb_flow_rotation/test_motion.py']
    write_json(OUT/'VALIDATION.json',dict(all_pass=True,tests_passed=count,scene=scene,motion_verification=motion,
        media=media,visual_inspection=inspection,flow_export_sha256=sha(OUT/'data/analytical_flow_si.vtu'),
        pptx=dict(slides=2,embedded_videos=1,sha256=sha(deck),native_powerpoint_UI_tested=False),
        source_sha256={p:sha(ROOT/p) for p in sources},manual_scientific_review='PENDING_USER_REVIEW'))
    (OUT/'REVIEW_ZH.md').write_text(f'''# RBC 与 MB 的水动力平移和旋转：共同可视化

## 本版结果

在同一血管和时间轴中展示正常双凹 RBC 与球形 MB 的**平移及旋转**。RBC 的三维姿态直接驱动双凹网格的刚体旋转；MB 的体坐标系随已计算角速度转动，白色条带和金色表面标记帮助观察球体旋转。所有标记都随同一个刚体变换，未单独指定动画转速。

保留此前黑底、透明管壁、英文标注和短尾迹的展示风格。正常红细胞不发生形变。主动画与 PPT 均为 1920 × 1080、16:9；视频 24 fps、18 秒，对应同一段 30 ms 物理过程。

## 流场信息和场景选择

此前确认的较宽理想化血管**有完整的解析速度信息**：内径 14 µm、平均流速 2 mm/s、中心线速度 4 mm/s 的充分发展 Poiseuille 背景流。可直接计算每个位置的速度梯度和涡量，因此满足本次平移、剪切转动和局部流体转动所需的数据条件。本轮继续使用该场景，无需转至原始分叉血管。

`data/analytical_flow_si.vtu` 导出速度、速度梯度、涡量与速度大小，几何采用 SI 单位。该文件是解析场的采样，不是新 FEM 求解结果；管壁附近显示网格为阶梯近似，运动与旋转始终采样解析场。固定观察窗口长度 40 µm，左右端为显示边界。

## 沿用的动力学模型

- **平移**：两类粒子沿用已有无额外外力的过阻尼随流运动；当前轴向不变的速度场允许精确计算中心轨迹。
- **RBC 旋转**：复用现有 `rbc_orientation.angular_velocity` 和 `advance_orientation`，由局部涡量、应变率和细胞朝向决定角速度。采用原有体积匹配扁球体的 Jeffery 旋转近似；正常双凹表面用于保持所需可视形态，未重新求解双凹物体的完整水动力阻力。
- **MB 旋转**：复用现有球形粒子角速度，即局部涡量的一半。姿态使用该角速度的精确恒定旋转解。表面标记是姿态指示，不表示膜上的真实结构或膜切向运动。

Jeffery 角速度及其涡量、应变率两项的依据可见 [Pujara 等，JFM 2021，式 1.1](https://doi.org/10.1017/jfm.2021.543)。本次沿用该形式检查现有实现，未将论文中的湍流研究场景代入本管流展示。

两类粒子受到**同一个给定背景流**驱动；没有粒子间水动力反馈、近场润滑或膜弹性求解。本版体现现有简化模型中的水动力平移和旋转，不是完整有限尺寸细胞悬浮液的双向耦合计算。

## 数量、尺寸与数值证据

共 9 个 RBC、28 个 MB；它们从上游流经固定观察窗口，不代表实验浓度。RBC 直径 {scene['original_geometry']['diameter_um']:.6f} µm，MB 直径 {scene['records'][-1]['radius_m']*2e6:.6f} µm。RBC 采用三种初始朝向；其后姿态均由上述模型积分产生。MB 两条初始通道选在能保证任意 RBC 朝向下几何间距的位置，未使用碰撞修正。

| 检查项目 | 结果 |
|---|---:|
| 永久测试及相关回归测试 | {count} 项通过 |
| RBC 姿态最大子步 | 5 µs |
| 10 / 5 µs 最大姿态差 | {motion['maximum_refinement_orientation_error_deg']:.6f}° |
| 与独立解析 Jeffery 短轴解的最大差 | {motion['maximum_exact_Jeffery_axis_error_deg']:.6f}° |
| 四元数最大模长误差 | {motion['maximum_quaternion_norm_error']:.3e} |
| RBC 角速度范围 | {motion['RBC_angular_speed_range_rad_s'][0]:.3f}–{motion['RBC_angular_speed_range_rad_s'][1]:.3f} rad/s |
| MB 角速度范围 | {motion['MB_angular_speed_range_rad_s'][0]:.3f}–{motion['MB_angular_speed_range_rad_s'][1]:.3f} rad/s |
| 任意朝向下管壁保守间隙 | {motion['minimum_all_orientation_wall_bound_um']:.6f} µm |
| 全时间包围球最小间隙 | {motion['minimum_all_time_bounding_sphere_pair_gap_um']:.6f} µm |

测试另以独立高精度 DOP853 四元数 ODE 核验 RBC 完整旋转；以旋转向量解析解核验 MB 姿态；通过有限差分核验流场梯度和涡量。全部 432 帧视频解码成功，每帧同时包含两类粒子。PPT 包结构与内嵌视频哈希核验通过，未执行原生 PowerPoint 界面播放测试。助手已检查解码关键帧，用户人工科学审核待进行。

## 文件入口

- `RBC_MB_Flow_Rotation_Presentation.pptx`：静态总览 + 内嵌动画，共 2 页。
- `animations/RBC_MB_flow_rotation.mp4`：主动画。
- `figures/RBC_MB_flow_rotation_overview.png`、`RBC_MB_flow_rotation_storyboard.png`：总览和四时刻截图。
- `figures/rotation_verification.png` / `.svg`：RBC 朝向变化与 MB 累积转角。
- `data/trajectories_and_rotation.csv`、`motion.npz`：同一时间轴的中心、速度、四元数和角速度。
- `data/SCENE.json`、`MOTION_VERIFICATION.json`、`VALIDATION.json`：流场、模型和验证记录。
- `OUTPUT_SHA256SUMS.txt`：输出校验清单。
''')
    (OUT/'OPEN_RESULTS.html').write_text('''<!doctype html><html lang="zh"><meta charset="utf-8"><title>RBC 与 MB：平移和旋转</title>
<style>body{max-width:1320px;margin:36px auto;padding:0 24px;background:#000;color:#eef3fa;font:18px system-ui}p{line-height:1.7}a{color:#67e3f6}video,img{width:100%;margin:14px 0}</style>
<h1>RBC 与 MB：共同流动与水动力旋转</h1><p>正常双凹红细胞与微泡在同一解析 Poiseuille 流场中运动和旋转。较宽理想化血管内径 14 µm，平均流速 2 mm/s；物理过程 30 ms，慢放为 18 秒。</p>
<p><a href="RBC_MB_Flow_Rotation_Presentation.pptx">PPT（内嵌动画）</a> · <a href="animations/RBC_MB_flow_rotation.mp4">MP4</a> · <a href="REVIEW_ZH.md">模型与验证说明</a> · <a href="data/trajectories_and_rotation.csv">轨迹与姿态 CSV</a> · <a href="data/analytical_flow_si.vtu">解析流场 VTU</a></p>
<video controls preload="metadata" poster="figures/RBC_MB_flow_rotation_overview.png" src="animations/RBC_MB_flow_rotation.mp4"></video>
<p>红色：正常 RBC；青色：MB。MB 表面标记用于显示球体旋转。采用现有 Jeffery / 半涡量角速度近似，不含粒子间流体反馈。</p>
<img alt="视频四个时刻" src="figures/RBC_MB_flow_rotation_storyboard.png"><img alt="旋转数值核验" src="figures/rotation_verification.png"></html>''')
    paths=sorted(p for p in OUT.rglob('*') if p.is_file() and p.name!='OUTPUT_SHA256SUMS.txt')
    (OUT/'OUTPUT_SHA256SUMS.txt').write_text(''.join(f'{sha(p)}  {p.relative_to(OUT)}\n' for p in paths))
    print(json.dumps(dict(all_pass=True,tests=count,pptx=str(deck),motion=motion,files=len(paths)),indent=2))


if __name__=='__main__':main()
