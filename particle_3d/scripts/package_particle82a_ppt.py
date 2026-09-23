#!/usr/bin/env python3
"""Package verified single-method presentation assets, with explicit scientific scope."""
from pathlib import Path
from collections import Counter
import hashlib,html,json,subprocess,zipfile,xml.etree.ElementTree as ET
from datetime import datetime,timezone
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'particle_3d/outputs/particle8_2a_ppt'
REPORT=ROOT/'particle_3d/reports/particle8_2a_ppt'


def read(p):return json.loads(Path(p).read_text())
def digest(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def write(p,value):Path(p).write_text(json.dumps(value,ensure_ascii=False,indent=2,allow_nan=False)+'\n')


def textbox(slide,x,y,w,h,text,size=22,color='EEF2F8'):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=shape.text_frame;tf.word_wrap=True
    for n,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if n==0 else tf.add_paragraph();p.text=line
        p.font.name='Microsoft YaHei';p.font.size=Pt(size);p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(12)
    return shape


def make_deck(scene,media,manifest):
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    def blank():
        s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(0,0,0)
        return s
    s=blank();s.shapes.add_picture(str(OUT/'figures/00_full_vessel_microbubble_tracks.png'),0,0,width=prs.slide_width,height=prs.slide_height)
    s.notes_slide.notes_text_frame.text='500 条独立有限尺寸微泡轨迹。使用固定入口位置、条件化粒径接纳方案。彩色为到达出口的路径，灰色为停止时保存的路径。颜色表示微泡速度；浅色球保持真实半径。完整血管图用于交代空间关系。'
    labels=['完整血管：旋转与微泡运动','分叉特写：微泡及轨迹','单微泡：保存路径的慢速回放']
    for i,record in enumerate(media['records']):
        s=blank();poster=OUT/'inspection'/Path(record['inspection_frames'][1]).name
        s.shapes.add_movie(str(OUT/'animations'/record['file']),0,0,prs.slide_width,prs.slide_height,
            poster_frame_image=str(poster),mime_type='video/mp4')
        note='前两段视频将独立轨迹的年龄错开叠加以便观察，不代表同一时刻的真实浓度或共同注入过程。' if i<2 else '这一段只播放一条保存的实际轨迹，画面底部显示自出生起的物理时间，播放时间经过拉长。为展示完整长路径，球在全长视图中较小；真实尺寸可结合分叉特写理解。'
        s.notes_slide.notes_text_frame.text=labels[i]+'。视频已内嵌，放映时点击画面播放。'+note+' 所有位置来自保存样本的分段线性插值，无延长路径，无球径放大。'
    s=blank();s.shapes.add_picture(str(OUT/'figures/02_multiview.png'),0,0,width=prs.slide_width,height=prs.slide_height)
    s.notes_slide.notes_text_frame.text='同一批 500 条路径的四个观察角度。旋转相机，物理坐标和血管比例不变。'
    s=blank();summary=scene['summary']
    textbox(s,.55,.35,12.2,.7,'500 条微泡轨迹｜展示数据说明',30)
    columns=[(.6,'500','已计算轨迹'),(4.9,str(summary['completed']),'到达出口'),(9.1,str(500-summary['completed']),'尚未到达出口')]
    for x,value,label in columns:
        textbox(s,x,1.3,3.6,.9,value,48,'77DCC9');textbox(s,x,2.24,3.6,.5,label,20,'BAC5D3')
    outlets=summary['outlets'];reasons=summary['end_reasons']
    body=(f"出口 01 / 02 / 03：{outlets['OUTLET_01']} / {outlets['OUTLET_02']} / {outlets['OUTLET_03']} 条\n"
          f"安全停止 {reasons.get('INTEGRATION_SAFETY_STOP',0)} 条；达到时间上限 {reasons.get('PHYSICAL_RESIDENCE_HORIZON_REACHED',0)} 条\n"
          '固定入口位置，粒径经局部可接纳条件筛选；不代表原始粒径总体\n'
          f"背景沿用原冻结流场：入口平均速度 {summary['frozen_inlet_mean_mm_s']:.3f} mm/s\n"
          '前两段为独立轨迹错时叠加；第三段显示单条轨迹的实际年龄')
    textbox(s,.6,3.05,12.1,3.8,body,22)
    s.notes_slide.notes_text_frame.text='本稿只展示一种方案，不包含方法对照。500 个 ID 按原始已接纳事件顺序选择，未按出口或完成与否挑选。安全停止及时间上限不能解释为生理性滞留。未改用独立的 2.0 mm/s 流场。'
    path=OUT/'Microbubble_500_Presentation.pptx';prs.save(path)
    # Verify embedded media bytes and the native package graph; no claim of Office UI execution.
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
        embedded={n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('ppt/media/') and n.endswith('.mp4')}
        assert sorted(embedded.values())==sorted(r['sha256'] for r in media['records'])
        slide_xml=[n for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')]
        assert len(slide_xml)==6
        for name in slide_xml:ET.fromstring(z.read(name))
    assert len(Presentation(path).slides)==6
    return dict(all_pass=True,file=path.name,sha256=digest(path),slides=6,embedded_videos=embedded,
        pptx_zip_integrity=True,embedded_video_bytes_match_decoded_MP4=True,
        native_powerpoint_UI_playback_tested=False,movie_playback='CLICK_TO_PLAY',
        slide_notes_in_Chinese=True)


def main():
    assert (OUT/'PPT_COMPLETE').exists()
    compute=read(OUT/'PPT_COMPUTE_VALIDATION.json');scene=read(OUT/'data/SCENE.json')
    media=read(OUT/'PPT_MEDIA_VALIDATION.json');audit=read(OUT/'PPT_REPLAY_AUDIT.json')
    sync=read(REPORT/'REMOTE_LOCAL_CHECKSUM_VERIFICATION.json');preserved=read(REPORT/'UPSTREAM_IMMUTABILITY_VERIFICATION.json')
    remote=read(OUT/'PPT_REMOTE_IMMUTABILITY.json');visual=read(REPORT/'AGENT_VISUAL_INSPECTION.json')
    manifest=read(OUT/'PPT_RENDER_MANIFEST.json');scope=read(REPORT/'USER_SCOPE_REVISION.json')
    for check in [compute,media,audit,sync,preserved,remote,visual]:assert check['all_pass']
    assert len(scene['records'])==500 and compute['quota']==500
    for r in scene['records']:
        assert digest(OUT/r['samples_path'])==r['samples_sha256']
        assert digest(OUT/r['metadata_path'])==r['metadata_sha256']
    for r in visual['inspected_files']:assert digest(OUT/r['file'])==r['sha256']
    tests=ET.parse(OUT/'PPT_REMOTE_TESTS.xml').getroot().findall('.//testcase')
    assert len(tests)==43 and all(c.find('failure') is None and c.find('error') is None and c.find('skipped') is None for c in tests)
    deck=make_deck(scene,media,manifest);write(OUT/'PPTX_VALIDATION.json',deck)
    summary=scene['summary'];commits=dict(Counter(r['computation_source_commit'] for r in scene['records']))
    validation=dict(AUTOMATED_CHECKS='PASS',PPT_SCOPE_RESULT='PASS',INTEGRATED_TRAJECTORIES=500,
        ALL_500_REACHED_OUTLET=summary['completed']==500,LARGE_15000_STAGE='NOT_COMPLETED_USER_REVISED_SCOPE',
        selected_method='B_FIXED_ANCHOR_CONDITIONAL_SIZE',global_best_physical_method_proven=False,
        selected_before_integration=True,summary=summary,tests_passed=43,replay_audit=audit,
        pptx=deck,manual_visual_review='PENDING_USER_REVIEW',agent_visual_review=visual,
        execution_source_commits=dict(compute=compute['source_commit'],render=remote['snapshots'][1]['source_commit'],
            trajectory_counts_by_source=commits,packaging=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip()),
        newly_integrated=500-len(compute['reused']),reused_exact_saved_trajectories=len(compute['reused']),
        remote_compute_host=compute['hostname'],compute_seconds=compute['seconds'],remote_local_sync=sync,
        upstream_immutability=preserved,remote_immutability=remote,
        display_bubble_selection=manifest['moving_selection'],display_bubble_pool=len(manifest['moving_display_ids']),
        no_MB_MB_interactions=True,no_population_concentration_claim=True,
        single_journey_id=manifest['single_journey_id'],single_journey_selection=manifest['single_selection'],
        packaging_source_sha256=digest(__file__),created_utc=datetime.now(timezone.utc).isoformat())
    write(REPORT/'PPT_VALIDATION.json',validation);write(OUT/'PPT_VALIDATION.json',validation)
    rows='\n'.join(f"| {k} | {v} |" for k,v in summary['end_reasons'].items())
    review=f'''# 500 条微泡轨迹：PPT 展示审核说明

本轮按用户最新要求停止原 15,000 条大规模积分任务，改为单方案、500 条轨迹的展示。原任务已经生成的数据保留，未宣称原大样本阶段完成。

## 可直接使用的文件

- `Microbubble_500_Presentation.pptx`：6 页、16:9、黑底，内嵌 3 段视频；放映时点击视频画面播放。
- `OPEN_RESULTS.html`：本地打开可浏览高清图和播放独立 MP4。
- `figures/`：4 张 3840 × 2160 PNG，以及一张 2880 × 1800 视频关键帧总览。
- `animations/`：3 段 1920 × 1080、24 fps、18 秒 H.264 MP4。
- `data/trajectory_catalog.csv`、`data/SCENE.json`、`data/B/trajectories/`：全部 500 条记录及原始保存样本。

## 本次展示的是哪一批微泡

采用现有方案 B：保留入口位置，在该位置的可接纳条件下筛选粒径。选择依据是已有入口覆盖证据和小规模试算表现，服务于这次展示；本轮没有证明某方案在物理上全局最优，也没有进行方法对照展示。按原共同事件顺序取前 500 个已接纳 ID，未按最终出口或完成状态挑选。其中复用 {len(compute['reused'])} 条完全一致的已存轨迹，新增积分 {500-len(compute['reused'])} 条；没有额外扩大配额。

平均直径 {summary['diameter_mean_um']:.4f} µm，范围 {summary['diameter_range_um'][0]:.4f}–{summary['diameter_range_um'][1]:.4f} µm。原始 SonoVue 分布均值为 {summary['original_sonovue_mean_um']:.4f} µm，因此这批是经过接纳条件筛选的尺寸总体，不能当作原始无条件总体。

背景沿用 P8.1 原冻结 FEM 流场，入口平均速度 **{summary['frozen_inlet_mean_mm_s']:.6f} mm/s**。本轮借用之前 2.0 mm/s 流场的可视化风格，没有切换背景流场。有限尺寸球形微泡模型、积分主体和边界处理保持原实现；不加入微泡间作用或红细胞耦合。

## 500 条计算结果

到达出口 {summary['completed']} 条：出口 01 为 {summary['outlets']['OUTLET_01']} 条，出口 02 为 {summary['outlets']['OUTLET_02']} 条，出口 03 为 {summary['outlets']['OUTLET_03']} 条。另有 {500-summary['completed']} 条尚未到达出口。

| 结束原因 | 数量 |
|---|---:|
{rows}

安全停止表示积分触发原有数值守卫；时间上限表示在设定观察期限内尚未完成。两者不能解释为生理性堵塞、微泡吸附或确定的最终去向。出口完成采用原有中心首次穿越出口判定。本轮实际分流仍主要集中在出口 02，不通过补线或改变路径制造均匀出口覆盖。

## 如何读图与动画

灰色半透明表面为原血管几何，浅色球按真实保存半径绘制；彩色线为已到达出口的轨迹，按保存的微泡速度着色，灰色细线为停止时保存的未完成路径。右侧统一色标为 0–{manifest['color_scale_mm_s'][1]:.2f} mm/s，样本最大微泡速度 {summary['speed_max_mm_s']:.6f} mm/s。色标不是新的背景流场速度求解。

第一段为完整血管旋转，第二段放大关键分叉，第三段为单微泡的完整保存路径回放。前两段全部 500 条路径可见，活动球从已完成轨迹中选取 {len(manifest['moving_display_ids'])} 个 ID，并错开各自轨迹年龄便于观察；这不是同一时刻 500 个微泡的真实共同注入过程，不能据此读取浓度、到达率或相互碰撞。所有显示位置仍严格来自该微泡保存样本的分段线性插值。

第三段选用已完成轨迹中保存路程最长的一条（ID {manifest['single_journey_id']}），用于说明长路径；它是明确选取的示例，不是代表性出口比例样本。底部年龄为该轨迹的真实物理时间，视频播放时间拉长。单球在完整长路径视图中较小，尺寸观察可结合第二段特写。

相机绕模型旋转；物理坐标、血管比例及微泡半径没有旋转重写、拉伸或放大。全血管视图每帧检查几何投影边界；特写视图有意裁切远端。未对未完成轨迹外推，也没有在出口后继续画球。

## 验证与来源

远程计算主机 `{compute['hostname']}`，使用 8 个 CPU worker，新增积分用时约 {compute['performance']['seconds']/60:.1f} 分钟。服务器为 Ryzen 7 7800X3D，容器 CPU 配额 7.68 核；本积分使用 CPU，并非 GPU 求解。

43 项永久测试通过；3 段视频的 {audit['verified_video_frames']} 帧全部成功解码；独立核验 {audit['verified_bubble_states']} 个逐帧微泡状态，最大重建位置误差 {audit['maximum_position_replay_error_m']:.3g} m。全部 500 条样本哈希与来源记录一致，完成轨迹的实际末段出口重新分类通过。

远程与 WSL 共 {sync['checked_files']} 个文件逐字节 SHA256 一致。原输入及历史结果 {preserved['checked_files']} 个文件重新哈希，无变化。远程计算、渲染的源代码快照及 Frozen / SonoVue 输入也已复核。

计算源代码 `{compute['source_commit']}`；渲染源代码 `{remote['snapshots'][1]['source_commit']}`。6 条复用轨迹保留各自原代码来源，详见 `PPT_VALIDATION.json`。打包脚本与测试文件保留在项目中。

代理已查看正式静态图和每段视频的首、中、末关键帧，检查黑底、文字、右侧色标、球和轨迹、相机视角。PPTX 包结构及 3 个内嵌视频的字节校验通过；当前 WSL 没有原生 PowerPoint，未宣称已在 PowerPoint 界面试播。独立 MP4 已逐帧解码，可单独插入演示稿。

`AUTOMATED_CHECKS = PASS`；`PPT_SCOPE_RESULT = PASS`；`MANUAL_VISUAL_REVIEW = PENDING_USER_REVIEW`。这里的 PASS 指本轮 500 条计算记录与展示交付完成，不表示 500 条均到达出口，也不替代用户的科学审核。
'''
    (REPORT/'PPT_REVIEW.md').write_text(review);(OUT/'PPT_REVIEW.md').write_text(review)
    titlemap={'00_full_vessel_microbubble_tracks':'完整血管与 500 条轨迹','01_branch_microbubble_detail':'分叉特写',
        '02_multiview':'四个观察角度','03_single_microbubble_journey':'单微泡保存路径','04_animation_storyboard':'动画关键帧总览'}
    blocks=[]
    for record in media['records']:
        name=html.escape(record['file']);blocks.append(f'<section><h2>{name}</h2><video controls preload="metadata" src="animations/{name}"></video><p><a href="animations/{name}" download>下载 MP4</a></p></section>')
    for p in sorted((OUT/'figures').glob('*.png')):
        blocks.append(f'<section><h2>{titlemap[p.stem]}</h2><a href="figures/{p.name}"><img src="figures/{p.name}" alt="{titlemap[p.stem]}"></a></section>')
    page='''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>500 条微泡轨迹 · PPT 素材</title><style>body{margin:32px auto;max-width:1180px;padding:0 22px;background:#050708;color:#edf3f8;font:18px/1.65 system-ui,sans-serif}a{color:#8cdbf0}h1{font-size:32px}h2{font-size:22px}section{margin-top:42px}video,img{display:block;width:100%;background:#000}p{color:#c1ccd7}.links{display:flex;gap:24px;flex-wrap:wrap}</style>
<h1>三维血管中的微泡与轨迹</h1><p>单方案 · 500 条独立轨迹 · 黑底展示素材</p>
<p class="links"><a href="Microbubble_500_Presentation.pptx" download>下载 6 页 PPT（内嵌视频）</a><a href="PPT_REVIEW.md">中文说明</a><a href="PPT_VALIDATION.json">验证记录</a><a href="data/trajectory_catalog.csv">500 条轨迹 CSV</a></p>
'''+f'<p>已到达出口 {summary["completed"]} 条；出口 01 / 02 / 03：{summary["outlets"]["OUTLET_01"]} / {summary["outlets"]["OUTLET_02"]} / {summary["outlets"]["OUTLET_03"]}。其余为安全停止或时间上限。</p>'+'''
<p>固定入口位置、条件化粒径接纳；背景为原冻结流场（入口均速 0.353 mm/s）。前两段活动球为独立轨迹错时叠加，第三段为单条实际路径年龄回放。颜色为微泡速度，球保持真实半径。</p>
'''+''.join(blocks)+'</html>'
    (OUT/'OPEN_RESULTS.html').write_text(page)
    bundle=OUT/'Microbubble_500_PPT_assets.zip'
    members=[OUT/'Microbubble_500_Presentation.pptx',OUT/'OPEN_RESULTS.html',OUT/'PPT_REVIEW.md',OUT/'PPT_VALIDATION.json',OUT/'PPTX_VALIDATION.json',OUT/'data/trajectory_catalog.csv']
    members+=sorted((OUT/'figures').glob('*.png'))+sorted((OUT/'animations').glob('*.mp4'))
    with zipfile.ZipFile(bundle,'w',zipfile.ZIP_DEFLATED) as z:
        for p in members:z.write(p,p.relative_to(OUT))
    with zipfile.ZipFile(bundle) as z:assert z.testzip() is None
    write(REPORT/'DELIVERY_MANIFEST.json',dict(all_pass=True,package=dict(file=str(bundle),sha256=digest(bundle)),
        files={str(p.relative_to(OUT)):digest(p) for p in members},raw_500_trajectory_data_available_at=str(OUT/'data'),
        source_script_sha256=digest(__file__)))
    print(json.dumps(dict(pptx=deck['file'],slides=6,zip=str(bundle),automated_checks='PASS'),indent=2))


if __name__=='__main__':main()
