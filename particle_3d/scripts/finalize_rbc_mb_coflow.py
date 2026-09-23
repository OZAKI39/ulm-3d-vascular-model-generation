"""Package the shared-clock RBC / MB flow visualization for presentation."""
import hashlib
import json
import xml.etree.ElementTree as ET
import zipfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

from prepare_rbc_mb_coflow import OUT, ROOT, sha, write_json


def main():
    read = lambda name: json.loads((OUT / name).read_text())
    scene = read('data/SCENE.json')
    media = read('MEDIA_VALIDATION.json')
    inspection = read('VISUAL_INSPECTION.json')
    suites = list(ET.parse(OUT / 'tests.xml').getroot().iter('testsuite'))
    count = sum(int(s.get('tests', 0)) for s in suites)
    assert count == 4
    assert all(int(s.get(k, 0)) == 0 for s in suites for k in ['failures', 'errors', 'skipped'])
    assert media['all_pass'] and inspection['all_pass']
    video = OUT / 'animations/RBC_MB_coflow.mp4'
    assert sha(video) == media['video_sha256']
    assert sha(OUT / 'data/coflow.npz') == media['source_npz_sha256'] == scene['array_sha256']
    assert sha(OUT / 'data/normal_rbc_si.vtp') == scene['shape_source_sha256']
    for relative, digest in inspection['inspected_image_sha256'].items():
        assert sha(OUT / relative) == digest

    deck = OUT / 'RBC_MB_Coflow_Presentation.pptx'
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    notes = (
        'RBC 和 MB 共同流动的同屏展示。沿用已确认的直径 14 µm 理想化直血管、'
        '平均速度 2 mm/s 的解析 Poiseuille 场以及正常双凹红细胞几何。'
        '同一时间轴为 0–30 ms，视频慢放为 18 秒。'
        '共保存 9 个 RBC 和 28 个 MB 的轨迹；40 µm 长的固定观察窗口内可见数量随时间变化。'
        '这些数量用于展示，未定义为实验浓度或血细胞比容。'
        '两类粒子沿用各自独立的随流运动模型；本片不表示粒子间相互作用求解。'
        '正常双凹形态和粒子物理尺寸全程保持，尾迹保留最近 2 ms。'
    )
    for movie in [False, True]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(0, 0, 0)
        if movie:
            slide.shapes.add_movie(str(video), 0, 0, prs.slide_width, prs.slide_height,
                poster_frame_image=str(OUT / 'figures/RBC_MB_coflow_overview.png'), mime_type='video/mp4')
        else:
            slide.shapes.add_picture(str(OUT / 'figures/RBC_MB_coflow_overview.png'), 0, 0,
                width=prs.slide_width, height=prs.slide_height)
        slide.notes_slide.notes_text_frame.text = notes + (' 放映时点击视频播放。' if movie else '')
    prs.save(deck)
    with zipfile.ZipFile(deck) as archive:
        assert archive.testzip() is None
        movies = [name for name in archive.namelist() if name.endswith('.mp4')]
        assert len(movies) == 1
        assert hashlib.sha256(archive.read(movies[0])).hexdigest() == media['video_sha256']
    assert len(Presentation(deck).slides) == 2

    sources = [
        'particle_3d/scripts/prepare_rbc_mb_coflow.py',
        'particle_3d/scripts/render_rbc_mb_coflow.py',
        'particle_3d/scripts/finalize_rbc_mb_coflow.py',
        'particle_3d/scripts/render_normal_rbc_encounter.py',
        'particle_3d/src/particle_3d/normal_rbc_encounter.py',
        'particle_3d/src/particle_3d/integrator.py',
        'particle_3d/src/particle_3d/rbc_integrator.py',
        'particle_3d/tests/rbc_mb_coflow/test_coflow.py',
    ]
    write_json(OUT / 'VALIDATION.json', dict(
        all_pass=True, deliverable='SHARED_CLOCK_RBC_MB_COFLOW_VISUALIZATION',
        tests_passed=count, scene=scene, media=media, visual_inspection=inspection,
        pptx=dict(slides=2, embedded_videos=1, sha256=sha(deck), native_powerpoint_UI_tested=False),
        source_sha256={p: sha(ROOT / p) for p in sources},
        manual_scientific_review='PENDING_USER_REVIEW',
    ))
    geometry = scene['original_geometry']
    report = f'''# RBC 与 MB 共同流动：展示说明

本版将正常双凹红细胞（红色）和微泡（青色）放在**同一血管、同一物理时间轴和同一视角**中，呈现两类粒子的共同流动与短尾迹。采用黑底、透明灰色管壁、英文标注及速度色标，保持此前血管流场展示风格。

## 可直接使用的文件

- `RBC_MB_Coflow_Presentation.pptx`：16:9、2 页，第一页静态总览，第二页内嵌视频。
- `animations/RBC_MB_coflow.mp4`：1920 × 1080，24 fps，18 秒，432 帧。
- `figures/RBC_MB_coflow_overview.png`：可直接插入 PPT 的总览图。
- `figures/RBC_MB_coflow_storyboard.png`：从视频直接解码的四个时刻，3840 × 2160。
- `OPEN_RESULTS.html`：本地预览和下载入口。

## 场景与运动

| 项目 | 数值 |
|---|---:|
| 理想化血管内径 | 14 µm |
| 固定观察窗口长度 | 40 µm |
| 流场平均 / 中心线速度 | 2.0 / 4.0 mm/s |
| 正常红细胞直径 | {geometry['diameter_um']:.6f} µm |
| 正常红细胞参考体积 | {geometry['volume_fL']:.6f} fL |
| 微泡直径 | {scene['records'][-1]['radius_m'] * 2e6:.6f} µm |
| 完整轨迹数量 | 9 个 RBC + 28 个 MB |
| 窗口内可见 RBC 数量 | {media['minimum_visible_counts']['RBC']}–{media['maximum_visible_counts']['RBC']} |
| 窗口内可见 MB 数量 | {media['minimum_visible_counts']['MB']}–{media['maximum_visible_counts']['MB']} |
| 共同物理时间 | 0–30 ms |
| 视频播放时长 | 18 s |
| 尾迹时间范围 | 最近 2 ms |

沿用用户此前确认的较宽理想化直血管。两类粒子在同一个解析 Poiseuille 背景流中独立随流运动，平移速度取各自中心处的局部流速；本场沿轴向不变，可以直接求得与已有粒子推进器一致的位置。正常双凹 RBC 的网格和尺寸直接复用先前几何文件。当前朝向下短轴保持不变，因此画面中保持正常双凹形态。

9 个 RBC 与 28 个 MB 是有限展示队列，初始位置覆盖观察窗口上游。窗口边缘代表局部观察范围，粒子从边缘进入或离开属于可视范围变化；管道在该理想化模型中继续延伸。尾迹表示实际中心位置经过的路径。粒子按物种着色；速度色标仅对应背景参考流线。部分粒子可能在二维投影中重叠，三维位置和尺寸检查确认它们保持分离。

这是两类粒子运动的**共同可视化**，未求解粒子间接触或流体耦合反馈；也没有使用本轮中止的耦合求解器原型。展示数量不用于推断浓度或血细胞比容。理想化场景已在每帧标明，并非原分叉 FEM 血管中的群体轨迹结果。

## 验证

4 项永久测试通过：闭式轨迹与已有 RBC / MB 推进器一致；共同时间轴、数组和原始几何校验通过；整个 30 ms 内具有保守的粒子间距与管壁余量；每个保存时刻均有两类粒子可见。

视频完整解码 432 帧且各帧不同；尺寸、帧率、两类粒子可见性和关键帧颜色检查通过。助手已检查总览及解码关键帧。PPT 包结构、页数和内嵌视频哈希检查通过，未执行原生 PowerPoint 界面播放测试。用户人工科学审核尚待进行。

轨迹见 `data/trajectories.csv` 和 `data/coflow.npz`；逐帧可见数量和物理时刻见 `data/replay_frames.json`；完整机器记录见 `VALIDATION.json`；文件哈希见 `OUTPUT_SHA256SUMS.txt`。
'''
    (OUT / 'REVIEW_ZH.md').write_text(report)
    (OUT / 'OPEN_RESULTS.html').write_text('''<!doctype html><html lang="zh"><meta charset="utf-8">
<title>RBC 与 MB 共同流动</title><style>body{max-width:1320px;margin:36px auto;padding:0 24px;background:#000;color:#eaf0fa;font:18px system-ui}p{line-height:1.7}a{color:#67e3f6}video,img{width:100%;margin:14px 0}</style>
<h1>RBC 与 MB 共同流动</h1><p>正常双凹红细胞与微泡同屏展示，采用此前确认的 14 µm 理想化血管。共同物理时间 30 ms，慢放为 18 秒。</p>
<p><a href="RBC_MB_Coflow_Presentation.pptx">PPT（内嵌视频）</a> · <a href="animations/RBC_MB_coflow.mp4">MP4</a> · <a href="figures/RBC_MB_coflow_overview.png">总览图</a> · <a href="REVIEW_ZH.md">中文说明</a> · <a href="VALIDATION.json">验证记录</a> · <a href="data/trajectories.csv">CSV 轨迹</a></p>
<video controls preload="metadata" poster="figures/RBC_MB_coflow_overview.png" src="animations/RBC_MB_coflow.mp4"></video>
<p>红色：正常红细胞；青色：微泡。尾迹保留最近 2 ms。两类粒子在相同背景流中独立运动，粒子数用于展示。</p>
<img alt="四个共同流动时刻" src="figures/RBC_MB_coflow_storyboard.png"></html>''')
    paths = sorted(p for p in OUT.rglob('*') if p.is_file() and p.name != 'OUTPUT_SHA256SUMS.txt')
    (OUT / 'OUTPUT_SHA256SUMS.txt').write_text(''.join(f'{sha(p)}  {p.relative_to(OUT)}\n' for p in paths))
    print(json.dumps(dict(all_pass=True, tests=count, pptx=str(deck), files=len(paths)), indent=2))


if __name__ == '__main__':
    main()
