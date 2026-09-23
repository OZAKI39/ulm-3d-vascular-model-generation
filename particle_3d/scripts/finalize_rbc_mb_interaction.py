"""Package verified saved interaction data and presentation media."""
from pathlib import Path
import json,hashlib,zipfile,xml.etree.ElementTree as ET
import numpy as np
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from simulate_rbc_mb_interaction import ROOT,OUT,CASE,sha
from particle_3d.particle3_cases import write_json

def main():
    read=lambda name:json.loads((OUT/name).read_text())
    fine=read('data/fine_states.json');coarse=read('data/coarse_states.json')
    summary=read('data/fine_summary.json');media=read('MEDIA_VALIDATION.json');init=read('data/initialization.json')
    tests=ET.parse(OUT/'tests.xml').getroot();suites=list(tests.iter('testsuite'))
    count=sum(int(s.get('tests','0')) for s in suites)
    assert count>0 and all(int(s.get(k,'0'))==0 for s in suites for k in ['failures','errors'])
    errors=[]
    for a,b in zip(coarse,fine[::2]):
        assert abs(a['time_s']-b['time_s'])<1e-15
        errors.extend(np.linalg.norm(np.array(x['center_m'])-y['center_m']) for x,y in zip(a['particles'],b['particles']))
    assert max(errors)<.01e-6
    interaction=[s for s in fine[1:] if s['projection']['contact_duration_s']>0]
    first=interaction[0];prev=fine[fine.index(first)-1]
    contact_time=prev['time_s']+first['projection']['hit_time_s']
    correction=max(np.linalg.norm(np.array(s['particles'][1]['velocity_m_s'])-s['projection']['mb_free_velocity_m_s']) for s in fine[1:])
    prism=np.array([[p['center_m'] for p in s['particles']] for s in fine])
    relative_change=float(np.linalg.norm((prism[-1,1]-prism[-1,0])-(prism[0,1]-prism[0,0])))
    prereq=read('VISUAL_INSPECTION.json');assert prereq['all_pass'] and media['all_pass']
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor(0,0,0)
    slide.shapes.add_picture(str(OUT/'figures/Interaction_overview.png'),0,0,width=prs.slide_width,height=prs.slide_height)
    note='本算例采用新的入口平均 2.0 mm/s FEM 流场，1 个微泡与 1 个红细胞进行局部接触。红细胞为固定轴向、保体积且满足面积预算的胶囊形近似。接触使用 P4 无摩擦运动学约束，并非红细胞膜力学或非球形润滑力。不是旧的 500 条独立轨迹整体加入红细胞后的重算，也不是完整血细胞悬浮液。'
    slide.notes_slide.notes_text_frame.text=note
    slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.shapes.add_movie(str(OUT/'animations/Microbubble_RBC_interaction.mp4'),
        0,0,prs.slide_width,prs.slide_height,poster_frame_image=str(OUT/'figures/decoded_frame_0144.png'),mime_type='video/mp4')
    slide.notes_slide.notes_text_frame.text=note+' 物理模拟时间 10 ms，视频播放 18 s。固定相机；曲面接触段使用解析子时间回放，未放大粒径或人为偏转轨迹。右侧色标对应左侧整血管背景流线；特写粒子及轨迹按物种着色。'
    deck=OUT/'Microbubble_RBC_Interaction_Presentation.pptx';prs.save(deck)
    with zipfile.ZipFile(deck) as z:
        assert z.testzip() is None
        videos=[n for n in z.namelist() if n.endswith('.mp4')]
        assert len(videos)==1 and hashlib.sha256(z.read(videos[0])).hexdigest()==media['sha256']
    assert len(Presentation(deck).slides)==2
    sourcefiles=[ROOT/'particle_3d/src/particle_3d/rbc_mb_encounter.py',
        ROOT/'particle_3d/scripts/simulate_rbc_mb_interaction.py',ROOT/'particle_3d/scripts/render_rbc_mb_interaction.py',
        ROOT/'particle_3d/scripts/finalize_rbc_mb_interaction.py',ROOT/'particle_3d/tests/rbc_mb_interaction/test_encounter.py']
    core=['particle4_motion.py','pair_geometry.py','particle_shapes.py','wall_gap.py','field.py','rbc_capillary_surrogate.py','kinematic_contact.py']
    sourcefiles += [ROOT/'particle_3d/src/particle_3d'/x for x in core]
    validation=dict(all_pass=True,tests_passed=count,physical_time_s=.01,video_duration_s=18.,
        counts={'RBC':1,'MB':1},first_contact_time_s=contact_time,
        max_coarse_fine_position_difference_um=float(max(errors)*1e6),
        max_MB_contact_velocity_correction_mm_s=float(correction*1000),relative_displacement_change_um=relative_change*1e6,
        minimum_wall_gap_um=summary['minimum_wall_gap_m']*1e6,
        minimum_pair_gap_m=summary['minimum_pair_gap_m'],negative_gap_within_float64_roundoff=True,
        initial_gap_um=init['initial_pair_gap_m']*1e6,source_field_case=str(CASE),
        source_npz_sha256=summary['source_npz_sha256'],source_wall_sha256=summary['source_wall_sha256'],
        model=summary['model'],physics_scope={'RBC_membrane_mechanics':False,'RBC_rotation':False,'nonspherical_lubrication':False,
            'two_way_CFD_coupling':False,'physiological_suspension':False,'concentration_or_hematocrit_claim':False},
        integration='EXACT_P4_FIXED_CAP_HEMISPHERE_CONTACT_WITH_FROZEN_FREE_VELOCITY_PER_STEP',
        all_accepted_intervals_have_continuous_wall_and_cap_regime_certificates=True,
        position_projection=False,media=media,visual_inspection=prereq,
        pptx={'slides':2,'embedded_videos':1,'sha256':sha(deck),'native_powerpoint_UI_tested':False},
        source_sha256={str(p.relative_to(ROOT)):sha(p) for p in sourcefiles},manual_scientific_review='PENDING_USER_REVIEW')
    write_json(OUT/'VALIDATION.json',validation)
    report=f'''# 微泡—红细胞局部相互作用：中文审核报告

## 可展示的结果

新增 1 个微泡与 1 个红细胞的局部相遇算例。采用新 FEM 流场（入口平均 2.0 mm/s）、原始血管几何和实际有限尺寸粒子。物理时间 **10 ms**，以 **18 秒**慢放呈现。左侧给出整血管位置，右侧给出固定相机特写；黑底、透明灰色血管、右侧速度色标与之前流场保持一致。红色为红细胞胶囊近似，青色为微泡；速度色标仅对应整血管背景流线。

- `animations/Microbubble_RBC_interaction.mp4`：1920 × 1080、24 fps、432 帧。
- `Microbubble_RBC_Interaction_Presentation.pptx`：2 页，视频已内嵌；未执行原生 PowerPoint 界面播放测试。
- `figures/Interaction_overview.png`、`figures/Interaction_storyboard.png`：总览与真实解码关键帧。
- `data/fine_trajectory.csv`、`data/fine_states.json`、`data/fine_ledger.json`：轨迹、接触状态与连续时间证书。

## 计算结果

- 初始真实表面间隙 {init['initial_pair_gap_m']*1e9:.1f} nm；首次接触约 **{contact_time*1000:.6f} ms**。
- 细步长 10 µs，复核步长 20 µs；最大位置差 **{max(errors)*1e6:.6f} µm**。
- 最大微泡接触速度修正 **{correction*1000:.6f} mm/s**，来自实际运动学接触解。
- 相对中心位移变化 **{relative_change*1e6:.6f} µm**。接触期两者共同移动并发生相对滑移；本时间窗口没有宣称完全绕过或完成红细胞通行。
- 最小管壁间隙 **{summary['minimum_wall_gap_m']*1e6:.6f} µm**。最小粒子间隙 {summary['minimum_pair_gap_m']:.3e} m，为浮点舍入量级，逐项小于原几何误差预算。
- 连续时间覆盖误差 {summary['physical_time_coverage_error_s']:.3e} s；**{count} 项永久测试通过**，全部 432 帧解码通过。

## 采用的模型及其适用范围

流体速度来自新场 `flow_arrays_si.npz` 的原始四面体 P1 插值。红细胞沿用 C57BL/6 V0 已保存分布中的样本（seed 2026092002，RBC ID 41310，原直径 {init['rbc_original_geometry']['provenance']['D_um']:.6f} µm，体积 {init['rbc_original_geometry']['provenance']['V_fL']:.6f} fL）。本展示固定选择半径 1.3 µm 的面积可行胶囊，圆柱段长 {init['rbc']['cylindrical_length_m']*1e6:.6f} µm，体积保持原值，面积不超过原扁椭球的面积预算。胶囊轴向和形状在这段局部计算中固定；没有计算动态膜变形。微泡半径 {init['mb']['radius_m']*1e6:.6f} µm，来自项目已有验证算例尺寸。

接触复用 P4 无摩擦、几何尺度运动学最小修正。因胶囊无转动自由度，本算例的接触位于半球端部时，其局部约束与项目已有的双球平移接触解析解相同。新增适配器只在半球区域证书成立时使用该解析积分；完整红细胞仍是胶囊，没有用等效球替代红细胞形状。原逐步 Euler 接触积分在极小间隙处产生细分抖动，日志保留为 `trial_profile.log`；最终采用解析接触积分，并以独立 ODE、原凸几何距离引擎、时间步长减半检查进行验证。原 P4/P6.5 公共实现没有改写。

每个子步冻结起点的实际 FEM 自由速度，精确求首次接触时刻，再沿接触曲面推进。每个接受区间用位移上界与原始管壁间隙证明整段不穿壁，用轴向余量证明没有离开半球适用区。显示时重算同一接受区间的解析部分时间解，避免用接触端点间的直线插值产生可视化穿透。没有位置推开、扩大血管、缩小细胞或人为偏转轨迹。

这是为展示相互作用选择的局部初值，不是从入口连续注入的细胞群体；不对应浓度、血细胞比容或分流概率。663 个初始几何候选只用于寻找可容纳且自然接近的位置，不是 663 条积分轨迹。探索性 20 ms 延长在接近管壁时触发局部模型守卫，保留 `extended_coarse.log`；最终展示窗口限制为经过验证的 10 ms，不把停止解释为生理堵塞。

本模型包含几何排斥与接触滑移；不包含红细胞膜力学、非球形润滑力、红细胞反馈改变流场、声学振荡或完整生理悬浮液。项目现有非球形润滑模型尚未冻结，因此没有调用仅适用于球体的 P6.5 阻力公式来冒充红细胞相互作用。

新动画使用 2.0 mm/s 新场；原 500 条独立微泡轨迹仍使用旧冻结场。本轮没有把那 500 条轨迹改写成含红细胞耦合的数据。

## 审核状态

自动检查通过；生成图片与视频解码关键帧已由助手检查。人工科学审核由用户继续进行。完整机器记录见 `VALIDATION.json`，各文件校验值见 `OUTPUT_SHA256SUMS.txt`。
'''
    (OUT/'REVIEW_ZH.md').write_text(report)
    (OUT/'OPEN_RESULTS.html').write_text('''<!doctype html><html lang="zh"><meta charset="utf-8"><title>微泡与红细胞相互作用</title>
<style>body{background:#000;color:#eaf0fa;font:18px system-ui;max-width:1300px;margin:40px auto;padding:0 24px}a{color:#6ae3f2}img,video{width:100%;margin:18px 0}p{line-height:1.7}</style>
<h1>微泡—红细胞局部相互作用</h1><p>新 FEM 流场，入口平均 2.0 mm/s。1 个微泡 + 1 个红细胞胶囊近似；物理时间 10 ms，播放 18 秒。采用 P4 无摩擦几何接触，红细胞形状固定。</p>
<p><a href="Microbubble_RBC_Interaction_Presentation.pptx">PPT（内嵌视频）</a> · <a href="REVIEW_ZH.md">中文审核报告</a> · <a href="VALIDATION.json">验证记录</a> · <a href="data/fine_trajectory.csv">CSV 轨迹数据</a></p>
<video controls preload="metadata" poster="figures/Interaction_overview.png" src="animations/Microbubble_RBC_interaction.mp4"></video>
<img src="figures/Interaction_storyboard.png"><p>红细胞为保体积、满足面积预算的固定胶囊形近似。本算例展示局部接触滑移，不代表完整膜变形、流场双向耦合或生理血细胞悬浮液。</p></html>''')
    files=[p for p in OUT.rglob('*') if p.is_file() and p.name!='OUTPUT_SHA256SUMS.txt']
    (OUT/'OUTPUT_SHA256SUMS.txt').write_text(''.join(f'{sha(p)}  {p.relative_to(OUT)}\n' for p in sorted(files)))
    print(json.dumps({'all_pass':True,'tests':count,'max_position_difference_um':max(errors)*1e6,'pptx':str(deck)},indent=2))

if __name__=='__main__':main()
