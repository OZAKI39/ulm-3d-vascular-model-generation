"""Package the user-selected NORMAL RBC animation and audited evidence."""
import json,zipfile,hashlib,xml.etree.ElementTree as ET
import numpy as np
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from simulate_normal_rbc_encounter import OUT,ROOT,sha
from particle_3d.particle3_cases import write_json

def main():
    read=lambda p:json.loads((OUT/p).read_text())
    summary=read('data/fine_summary.json');geometry=read('data/GEOMETRY.json');media=read('MEDIA_VALIDATION.json')
    fine=np.load(OUT/'data/fine.npz');coarse=np.load(OUT/'data/coarse.npz')
    error=float(np.linalg.norm(fine['positions_m'][::2]-coarse['positions_m'],axis=2).max())
    assert error<.005e-6
    tree=ET.parse(OUT/'tests.xml');suites=list(tree.getroot().iter('testsuite'))
    count=sum(int(s.get('tests','0')) for s in suites)
    assert count==19 and all(int(s.get(k,'0'))==0 for s in suites for k in ['failures','errors','skipped'])
    assert media['all_pass'] and sha(OUT/'data/fine.npz')==media['source_trajectory_sha256']
    assert sha(OUT/'animations/Normal_RBC_microbubble_interaction.mp4')==media['video_sha256']
    assert sha(OUT/'data/normal_rbc_si.vtp')==media['geometry_sha256']
    records=read('data/replay_frames.json')
    d=np.array([25.,-55.,80.]);d/=np.linalg.norm(d);up=np.array([0.,.8,.55]);up/=np.linalg.norm(up);right=np.cross(up,d)
    radii=np.array([summary['rbc_diameter_m']/2,summary['mb_radius_m']])*1e6
    projected_max=0.
    for r in records:
        xyz=np.array(r['positions_m'])*1e6-np.array(r['camera_focal_point_um'])
        mx=(np.abs(xyz@right)+radii)/(10.5*(.84*1920)/(.72*1080))
        my=(np.abs(xyz@up)+radii)/10.5
        projected_max=max(projected_max,float(mx.max()),float(my.max()))
    assert projected_max<.9
    inspection=read('VISUAL_INSPECTION.json');assert inspection['all_pass']
    deck=OUT/'Normal_RBC_Interaction_Presentation.pptx'
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    notes='按用户确认，采用直径 14 µm 的较宽理想化直血管。红细胞是正常双凹圆盘形，整个过程保持原始形状，未采用胶囊形。背景为平均 2 mm/s 的解析 Poiseuille 流，不是原 FEM 场。几何接触采用固定朝向、赤道平面内无摩擦运动学模型，不代表完整血细胞膜力学或生理悬浮液。'
    slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor(0,0,0)
    slide.shapes.add_picture(str(OUT/'figures/Normal_RBC_interaction_overview.png'),0,0,width=prs.slide_width,height=prs.slide_height)
    slide.notes_slide.notes_text_frame.text=notes
    slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.shapes.add_movie(str(OUT/'animations/Normal_RBC_microbubble_interaction.mp4'),
        0,0,prs.slide_width,prs.slide_height,poster_frame_image=str(OUT/'figures/decoded_0100.png'),mime_type='video/mp4')
    slide.notes_slide.notes_text_frame.text=notes+' 模拟物理时间 30 ms，慢放为 18 秒；相机跟随两粒子，红细胞不发生转动或形变。轨迹尾迹保留最近 3 ms，位置与曲面接触回放均来自实际计算。放映时点击视频播放。'
    prs.save(deck)
    with zipfile.ZipFile(deck) as z:
        assert z.testzip() is None
        videos=[n for n in z.namelist() if n.endswith('.mp4')]
        assert len(videos)==1 and hashlib.sha256(z.read(videos[0])).hexdigest()==media['video_sha256']
    assert len(Presentation(deck).slides)==2
    sources=[ROOT/'particle_3d/src/particle_3d/normal_rbc_encounter.py',ROOT/'particle_3d/src/particle_3d/particle4_motion.py',
        ROOT/'particle_3d/scripts/simulate_normal_rbc_encounter.py',ROOT/'particle_3d/scripts/render_normal_rbc_encounter.py',
        ROOT/'particle_3d/scripts/finalize_normal_rbc_encounter.py',ROOT/'particle_3d/tests/rbc_mb_normal/test_normal_encounter.py']
    validation=dict(all_pass=True,user_selected_scene='WIDER_IDEALIZED_VESSEL_WITH_NORMAL_BICONCAVE_RBC',
        red_cell_count=1,microbubble_count=1,geometry=geometry,simulation=summary,
        tests_passed=count,maximum_timestep_refinement_error_um=error*1e6,
        all_432_frames_keep_both_particles_inside_view=True,maximum_normalized_projection=projected_max,
        media=media,visual_inspection=inspection,pptx=dict(slides=2,embedded_videos=1,sha256=sha(deck),native_powerpoint_UI_tested=False),
        source_sha256={str(p.relative_to(ROOT)):sha(p) for p in sources},manual_scientific_review='PENDING_USER_REVIEW')
    write_json(OUT/'VALIDATION.json',validation)
    report=f'''# 正常双凹红细胞—微泡相互作用：展示与审核说明

## 本轮最终采用的内容

按用户确认，采用**较宽的理想化血管**展示**正常双凹圆盘状红细胞**与微泡的相互作用，并在每帧标明 `IDEALIZED VESSEL`。已取消胶囊形红细胞展示。最终文件全部位于本目录 `rbc_mb_normal`。

动画呈现接近、沿侧缘接触滑移、分离三个阶段。黑底、透明灰色血管、速度色标和英文标注与之前流场图风格一致；红色是正常红细胞，青色是微泡。背景流线用速度着色，红细胞和微泡的轨迹按物种着色。保留最近 3 ms 的尾迹；相机随两粒子的中间位置平移，观察方向固定。

- `animations/Normal_RBC_microbubble_interaction.mp4`：1920 × 1080，24 fps，18 秒，432 帧。
- `Normal_RBC_Interaction_Presentation.pptx`：2 页、16:9、内嵌视频；包结构和视频字节校验通过，未执行原生 PowerPoint 界面播放测试。
- `figures/Normal_RBC_interaction_overview.png`：接触阶段总览。
- `figures/Normal_RBC_interaction_storyboard.png`：4 个直接解码关键帧，3840 × 2160。
- `data/fine_trajectory.csv`、`fine.npz`、`fine_intervals.json`：机器可读轨迹及逐区间证据。

## 尺寸与时间

| 项目 | 数值 |
|---|---:|
| 理想化血管内径 | 14 µm |
| 解析流场平均速度 | 2.0 mm/s |
| 解析流场中心线速度 | 4.0 mm/s |
| 红细胞直径 | {geometry['diameter_um']:.6f} µm |
| 红细胞参考体积 | {geometry['volume_fL']:.6f} fL |
| 中央厚度 | {geometry['center_thickness_um']:.6f} µm |
| 最大厚度 | {geometry['maximum_thickness_um']:.6f} µm |
| 微泡直径 | {summary['mb_radius_m']*2e6:.6f} µm |
| 真实模拟时间 | 30 ms |
| 播放时长 | 18 s |
| 首次接触 | {summary['first_contact_time_s']*1000:.6f} ms |
| 接触结束 | {summary['last_contact_time_s']*1000:.6f} ms |
| 最终表面间隙 | {summary['final_gap_m']*1e6:.6f} µm |

正常形态在整个动画中保持不变：只平移，不拉长、不缩小、不变为胶囊，也不动态旋转。两粒子的几何比例保持真实计算比例。

## 模型来源和范围

正常双凹表面采用 Evans–Fung 型参数轮廓，系数来自 [Dao、Lim 与 Suresh (2003)，式 7](https://www.mit.edu/~mingdao/papers/JMPS_2003_Red_Blood_Cell.pdf)。该论文给出人红细胞的双凹几何表达，本演示借用其无量纲形态作为**几何建模近似**；直径和体积沿用项目 C57BL/6 V0 几何样本（seed 2026092002，RBC ID 41310）。初始厚度系数通过体积约束一次确定，之后保持不变；这不是随运动发生的形变，也不是已经验证的小鼠膜形态测量模型。闭合网格体积与参考体积的相对差为 {geometry['mesh_relative_volume_error']:.3e}，源于显示网格离散。

血管是独立的 14 µm 理想化直圆管，背景采用平均 2 mm/s 的解析 Poiseuille 速度场。**没有把原 FEM 血管放大后伪装成原模型，也没有把这个场景当作原 FEM 结果。**此前 500 条微泡轨迹和新 FEM 流线均不被改写。

接触使用项目 P4 的无摩擦运动学最小速度修正。本演示限制红细胞朝向固定，微泡球心和红细胞中心都在双凹盘的赤道平面内。对处于圆盘外侧的微泡，真实双凹表面最近点位于圆形侧缘，因此该受限情形下的间隙和接触法向可以精确求出。复用 P4 已有的双球平移接触解析积分，仅用于这一圆形侧缘约束；红细胞的可视和接触实体仍是正常双凹形状，并非把红细胞换成球。永久测试另用真实双凹三角网格最近点距离验证这一对应关系。

每个 10 µs 子步使用当前位置解析场的自由速度，精确处理首次接触和接触曲面上的滑移。整个区间以运动上界和管壁余量证明不穿壁。显示时回放同一解析解的部分时间，避免直线跨越接触曲面。没有位置推开或人为指定分离路径。20 µs 步长复核后的最大中心位置差为 **{error*1e6:.6f} µm**。

这是**固定正常形态、固定朝向、赤道平面内的简化几何接触演示**；不是一般朝向下完整的三维细胞动力学，不计算膜弹性、非球形润滑力、转动或细胞反馈流场，也不代表真实血细胞比容、浓度或群体统计。图中“接触”表示运动学排斥约束，不表示已经求得真实膜接触力。

## 验证结果

**{count} 项永久测试通过**：包含正常双凹形状及体积、闭合网格、独立网格距离、独立接触 ODE、连续管壁余量、接触后分离、步长减半和解析回放；并运行现有 P4 接触投影回归测试。

最小管壁保守余量 {summary['minimum_wall_bound_m']*1e6:.6f} µm；最小粒子间隙 {summary['minimum_pair_gap_m']:.3e} m，处于浮点误差预算内。全部 432 帧解码通过且各帧不同；所有帧中两个粒子的包围体投影均位于视窗内。助手已检查总览及视频四个关键帧，双凹形态、接近/滑移/分离和理想化场景标注清楚。用户人工科学审核待进行。

完整机器记录见 `VALIDATION.json`；文件校验清单见 `OUTPUT_SHA256SUMS.txt`。
'''
    (OUT/'REVIEW_ZH.md').write_text(report)
    (OUT/'OPEN_RESULTS.html').write_text('''<!doctype html><html lang="zh"><meta charset="utf-8"><title>正常红细胞与微泡相互作用</title>
<style>body{max-width:1320px;margin:40px auto;padding:0 25px;background:#000;color:#eaf0fa;font:18px system-ui}p{line-height:1.7}a{color:#67e3f6}video,img{width:100%;margin:15px 0}</style>
<h1>正常双凹红细胞与微泡相互作用</h1><p>按确认采用 14 µm 理想化血管。正常双凹红细胞全程保持固定形态，演示微泡的接近、侧缘滑移和分离。物理时间 30 ms，视频慢放为 18 秒。</p>
<p><a href="Normal_RBC_Interaction_Presentation.pptx">PPT（内嵌视频）</a> · <a href="animations/Normal_RBC_microbubble_interaction.mp4">MP4</a> · <a href="REVIEW_ZH.md">中文审核报告</a> · <a href="VALIDATION.json">验证记录</a> · <a href="data/fine_trajectory.csv">CSV 轨迹</a></p>
<video controls preload="metadata" poster="figures/Normal_RBC_interaction_overview.png" src="animations/Normal_RBC_microbubble_interaction.mp4"></video>
<img src="figures/Normal_RBC_interaction_storyboard.png"><p>背景是解析 Poiseuille 流场。采用固定朝向和赤道平面内的无摩擦几何接触近似；本演示不代表原 FEM 几何、完整膜力学或血细胞悬浮液。</p></html>''')
    paths=[p for p in OUT.rglob('*') if p.is_file() and p.name!='OUTPUT_SHA256SUMS.txt']
    (OUT/'OUTPUT_SHA256SUMS.txt').write_text(''.join(f'{sha(p)}  {p.relative_to(OUT)}\n' for p in sorted(paths)))
    print(json.dumps(dict(all_pass=True,tests=count,pptx=str(deck),max_dt_error_um=error*1e6,files=len(paths)),indent=2))

if __name__=='__main__':main()
